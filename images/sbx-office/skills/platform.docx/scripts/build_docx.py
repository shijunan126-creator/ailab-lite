#!/usr/bin/env python3
"""Build a clean .docx from a small JSON spec. No network, no external templates."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _para(doc, text: str, *, size: int = 10.5, space_after: int = 8, align=None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.5
    if align is not None:
        paragraph.alignment = align
    _font(paragraph.add_run(str(text)), size=size)


def _table(doc, columns, rows) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    for index, value in enumerate(columns):
        header[index].text = ""
        _font(header[index].paragraphs[0].add_run(str(value)), size=9.5, bold=True)
        header[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values[: len(columns)]):
            cells[index].text = ""
            _font(cells[index].paragraphs[0].add_run(str(value)), size=9.5)
    doc.add_paragraph()


def build(spec: dict, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")

    if spec.get("title"):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(36)
        _font(paragraph.add_run(str(spec["title"])), size=22, bold=True, color="1F2937")
    if spec.get("subtitle"):
        _para(doc, spec["subtitle"], size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    for key in ("author", "date"):
        if spec.get(key):
            _para(doc, spec[key], size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for item in spec.get("sections") or []:
        if item.get("page_break_before"):
            doc.add_page_break()
        heading = str(item.get("heading") or "").strip()
        if heading:
            level = int(item.get("level") or 1)
            level = min(3, max(1, level))
            paragraph = doc.add_heading(level=level)
            _font(paragraph.add_run(heading), size={1: 16, 2: 13.5, 3: 11.5}[level], bold=True, color="1F4E79")
        for text in item.get("paragraphs") or []:
            _para(doc, text)
        for text in item.get("bullets") or []:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(4)
            _font(paragraph.add_run(str(text)), size=10.5)
        table = item.get("table")
        if table and table.get("columns"):
            _table(doc, table["columns"], table.get("rows") or [])
        image = item.get("image")
        if image:
            path = Path(str(image))
            if not path.is_file():
                raise SystemExit(f"image not found: {path}")
            doc.add_picture(str(path), width=Inches(5.9))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if item.get("caption"):
            _para(doc, item["caption"], size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("spec", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    if not isinstance(spec, dict):
        raise SystemExit("spec must be a JSON object")
    build(spec, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
