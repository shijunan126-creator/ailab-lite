"""Check basic document structure; rendering is needed for actual pagination."""

import sys

from docx import Document
from docx.oxml.ns import qn


def main(path):
    doc = Document(path)
    findings = []
    if (
        not any(p.text.strip() for p in doc.paragraphs)
        and not doc.tables
        and not doc.inline_shapes
    ):
        findings.append("document has no visible content")
    widths = [
        section.page_width - section.left_margin - section.right_margin
        for section in doc.sections
    ]
    if not widths or min(widths) <= 0:
        findings.append("invalid page margins")
    else:
        for number, shape in enumerate(doc.inline_shapes, 1):
            if shape.width > max(widths):
                findings.append(f"inline image {number} exceeds page width")
        for number, table in enumerate(doc.tables, 1):
            if sum(column.width or 0 for column in table.columns) > max(widths):
                findings.append(f"table {number} may exceed page width")
    for image in doc.part.element.xpath('.//*[local-name()="blip"]'):
        reference = image.get(qn("r:embed"))
        if reference and reference not in doc.part.rels:
            findings.append(f"missing image relationship {reference}")
    for finding in findings:
        print("WARN", finding)
    print(f"Checked document; {len(findings)} findings. Inspect pagination visually.")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: check.py document.docx")
    try:
        raise SystemExit(main(sys.argv[1]))
    except Exception as exc:
        print(f"ERROR check failed: {exc}", file=sys.stderr)
        raise SystemExit(2) from exc
