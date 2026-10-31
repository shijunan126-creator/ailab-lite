"""Report likely slide geometry problems; visual inspection remains necessary."""

import math
import sys
import unicodedata

from pptx import Presentation
from pptx.util import Pt


def text_units(text):
    return sum(1 if unicodedata.east_asian_width(ch) in "WF" else 0.55 for ch in text)


def main(path):
    deck = Presentation(path)
    findings = []
    for number, slide in enumerate(deck.slides, 1):
        text_shapes = []
        shapes = []
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > deck.slide_width
                or shape.top + shape.height > deck.slide_height
            ):
                findings.append(f"slide {number}: off-slide shape {shape.name}")
            if shape.is_placeholder and shape.has_text_frame and not shape.text.strip():
                findings.append(f"slide {number}: empty placeholder {shape.name}")
            if shape.width > 0 and shape.height > 0:
                shapes.append(shape)
            if (
                not shape.has_text_frame
                or not shape.text.strip()
                or shape.width <= 0
                or shape.height <= 0
            ):
                continue
            frame = shape.text_frame
            text_shapes.append(shape)
            available_width = max(
                1, shape.width - frame.margin_left - frame.margin_right
            )
            available_height = max(
                1, shape.height - frame.margin_top - frame.margin_bottom
            )
            estimated_height = 0
            for paragraph in frame.paragraphs:
                sizes = [run.font.size for run in paragraph.runs if run.font.size]
                font = max(sizes, default=paragraph.font.size or Pt(18))
                per_line = max(1, available_width / font)
                lines = sum(
                    max(1, math.ceil(text_units(line) / per_line))
                    for line in paragraph.text.split("\n")
                )
                estimated_height += lines * font * 1.1
            if estimated_height > available_height * 1.1:
                findings.append(
                    f"slide {number}: possible text overflow in {shape.name}"
                )
        for i, left in enumerate(text_shapes):
            for right in text_shapes[i + 1 :]:
                width = min(left.left + left.width, right.left + right.width) - max(
                    left.left, right.left
                )
                height = min(left.top + left.height, right.top + right.height) - max(
                    left.top, right.top
                )
                if (
                    width > 0
                    and height > 0
                    and width * height
                    > 0.15 * min(left.width * left.height, right.width * right.height)
                ):
                    findings.append(
                        f"slide {number}: overlapping text shapes {left.name} / {right.name}"
                    )
        if len(shapes) > 500:
            findings.append(
                f"slide {number}: only first 500 shapes checked for overlap"
            )
        for i, left in enumerate(shapes[:500]):
            for right in shapes[i + 1 : 500]:
                if left in text_shapes and right in text_shapes:
                    continue
                width = min(left.left + left.width, right.left + right.width) - max(
                    left.left, right.left
                )
                height = min(left.top + left.height, right.top + right.height) - max(
                    left.top, right.top
                )
                if width <= 0 or height <= 0:
                    continue
                area = width * height
                smaller = min(left.width * left.height, right.width * right.height)
                if 0.15 * smaller < area < 0.95 * smaller:
                    findings.append(
                        f"slide {number}: possible overlapping shapes {left.name} / {right.name}"
                    )
    for finding in findings:
        print("WARN", finding)
    print(
        f"Checked {len(deck.slides)} slides; {len(findings)} findings. Inspect layout visually."
    )
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: check.py presentation.pptx")
    try:
        raise SystemExit(main(sys.argv[1]))
    except Exception as exc:
        print(f"ERROR check failed: {exc}", file=sys.stderr)
        raise SystemExit(2) from exc
