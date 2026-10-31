"""Recalculate a copy and report spreadsheet errors and broken chart sources."""

import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ERRORS = {"#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}


def chart_reference(workbook, formula):
    if not formula or "#REF!" in formula or "[" in formula or "!" not in formula:
        return False
    sheet, cells = formula.rsplit("!", 1)
    sheet = sheet.strip("'").replace("''", "'")
    if sheet not in workbook.sheetnames:
        return False
    try:
        left, top, right, bottom = range_boundaries(cells.replace("$", ""))
    except ValueError:
        return False
    if (
        not all(isinstance(n, int) for n in (left, top, right, bottom))
        or (right - left + 1) * (bottom - top + 1) > 100000
    ):
        return False
    ws = workbook[sheet]
    return any(
        cell.value is not None
        for row in ws.iter_rows(
            min_row=top, max_row=bottom, min_col=left, max_col=right
        )
        for cell in row
    )


def main(source):
    source = Path(source).resolve()
    problems = []
    original = load_workbook(source, data_only=False)
    for sheet in original:
        for number, chart in enumerate(sheet._charts, 1):
            series = chart.series
            if not series:
                problems.append(f"{sheet.title}: chart {number} has no series")
            for item in series:
                references = []
                for field in ("val", "xVal", "yVal", "cat"):
                    data = getattr(item, field, None)
                    if data:
                        for ref in (
                            getattr(data, "numRef", None),
                            getattr(data, "strRef", None),
                        ):
                            if ref:
                                references.append(ref.f)
                if not references or any(
                    not chart_reference(original, formula) for formula in references
                ):
                    problems.append(
                        f"{sheet.title}: chart {number} has an empty or broken data reference"
                    )
    with tempfile.TemporaryDirectory(prefix="xlsx-check-") as tmp:
        root = Path(tmp)
        profile = root / "profile"
        result = subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "xlsx:Calc MS Excel 2007 XML",
                "--outdir",
                str(root),
                str(source),
            ],
            capture_output=True,
            timeout=90,
            check=True,
        )
        recalculated = root / (source.stem + ".xlsx")
        if not recalculated.is_file():
            raise RuntimeError(
                "LibreOffice produced no recalculated workbook: "
                + result.stderr.decode(errors="replace")[-300:]
            )
        values = load_workbook(recalculated, data_only=True, read_only=True)
        for sheet in values:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e" or (
                        isinstance(cell.value, str) and cell.value in ERRORS
                    ):
                        problems.append(
                            f"{sheet.title}!{cell.coordinate}: {cell.value}"
                        )
        values.close()
    original.close()
    for problem in problems:
        print("ERROR", problem)
    if not problems:
        print("OK: recalculated formulas and chart references")
    return bool(problems)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: check.py workbook.xlsx")
    try:
        raise SystemExit(1 if main(sys.argv[1]) else 0)
    except Exception as exc:
        print(f"ERROR check failed: {exc}", file=sys.stderr)
        raise SystemExit(2) from exc
