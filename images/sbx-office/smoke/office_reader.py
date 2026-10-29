"""Offline image contract: read four formats and verify an edited workspace file."""

import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

with tempfile.TemporaryDirectory(prefix="office-reader-", dir="/workspace") as folder:
    root = Path(folder)
    document = Document()
    document.add_paragraph("reader-before")
    document.save(root / "sample.docx")
    workbook = Workbook()
    workbook.active.append(["reader-sheet", 120])
    workbook.save(root / "sample.xlsx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "reader-slide"
    deck.save(root / "sample.pptx")
    subprocess.run(
        [
            "libreoffice",
            "-env:UserInstallation=file://" + str(root / "lo"),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(root),
            str(root / "sample.docx"),
        ],
        check=True,
        timeout=60,
    )
    for kind, expected in [
        ("docx", "reader-before"),
        ("xlsx", "reader-sheet"),
        ("pptx", "reader-slide"),
        ("pdf", "reader-before"),
    ]:
        output = root / kind
        subprocess.run(
            ["office-read", str(root / ("sample." + kind)), str(output)],
            check=True,
            timeout=70,
        )
        manifest = json.loads((output / "manifest.json").read_text())
        assert manifest["status"] == "ready", manifest
        assert expected in (output / "text.txt").read_text(), kind
    empty = root / "already-empty"
    empty.mkdir()
    assert (
        subprocess.run(
            ["office-read", str(root / "sample.docx"), str(empty)], check=False
        ).returncode
        == 2
    )
    racers = [
        subprocess.Popen(
            ["office-read", str(root / "sample.docx"), str(root / "race")],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        for _ in range(2)
    ]
    for process in racers:
        process.communicate(timeout=70)
    assert sorted(process.returncode for process in racers) == [0, 2]
    document.paragraphs[0].text = "reader-after"
    document.save(root / "sample.docx")
    subprocess.run(
        ["office-read", str(root / "sample.docx"), str(root / "edited")],
        check=True,
        timeout=70,
    )
    assert "reader-after" in (root / "edited/text.txt").read_text()
    assert "reader-before" not in (root / "edited/text.txt").read_text()
    assert (
        subprocess.run(
            ["office-read", str(root / "sample.docx"), str(root / "edited")],
            check=False,
        ).returncode
        != 0
    )
print("OFFICE READER PASS")
