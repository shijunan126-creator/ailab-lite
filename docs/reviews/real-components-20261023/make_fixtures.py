from pathlib import Path
from pptx import Presentation
from PIL import Image
import io,random,zipfile
root=Path('/Users/michael/.cache/ailab-real-verify-20261023/fixtures');root.mkdir(exist_ok=True)
with (root/'50m.bin').open('wb') as f:
 for _ in range(50):f.write(b'clean-upload\n'*(1024*1024//13)+b'x'*(1024*1024%13))
assert (root/'50m.bin').stat().st_size==50*1024**2
p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);im=io.BytesIO();Image.frombytes('RGB',(2600,2600),random.Random(42).randbytes(2600*2600*3)).save(im,'PNG');im.seek(0);s.shapes.add_picture(im,0,0);target=root/'20m.pptx';p.save(target)
name='padding.bin';n=20*1024**2-target.stat().st_size-76-2*len(name)
with zipfile.ZipFile(target,'a',zipfile.ZIP_STORED) as z:z.writestr(name,b'x'*n)
assert target.stat().st_size==20*1024**2
print('fixture bytes',[(p.name,p.stat().st_size) for p in root.iterdir()])
