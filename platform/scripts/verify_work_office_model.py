"""Opt-in metered model + real offline container check; never connects to the business DB.

Run with --execute, an explicit key env/catalog and an already running disposable container.
Uses the Work prompt and fixed tool schemas; this is not a five-service E2E.
"""

import argparse
import json
import os
import subprocess
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--key-env", type=Path, required=True)
    parser.add_argument("--catalog", type=Path, required=True)
    parser.add_argument("--docker-wrapper", type=Path, required=True)
    parser.add_argument("--container", required=True)
    args = parser.parse_args()
    if not args.execute:
        parser.error("--execute is required for metered calls")
    from dotenv import dotenv_values

    for key, value in dotenv_values(args.key_env).items():
        if key.startswith("AILAB_KEY_") and value:
            os.environ[key] = value
    import httpx

    from ailab_platform.agent import discovery, prompts, tools
    from ailab_platform.catalog import load_catalog
    from ailab_platform.internal_skills import load as internal_load

    catalog = load_catalog(args.catalog)
    spec = catalog.get(catalog.default_id())
    if not spec or not spec.available or spec.provider != "openai":
        raise RuntimeError("real model unavailable")
    allowed = {"exec", "read_file", "write_file", "load_skill", "update_plan", "deliver_file"}
    tool_catalog = [discovery.SEARCH, *[t for t in tools.tools_for_turn() if t["function"]["name"] in allowed]]
    messages = [
        {"role": "system", "content": prompts.system_prompt()},
        {
            "role": "user",
            "content": "请检查/workspace/inputs/report.docx和/workspace/inputs/budget.xlsx，"
            "把报告里的收入和表格North行收入从120改成175，其余文字、数据和表格公式保留。"
            "成品分别保存到/workspace/outputs/report.docx和budget.xlsx。修改前了解原文，修改后实际检查两份成品，简述结果。",
        },
    ]
    commands = []
    usage = {"prompt_tokens": 0, "completion_tokens": 0}

    def docker(*parts):
        result = subprocess.run(
            ["bash", str(args.docker_wrapper), "exec", args.container, *parts],
            capture_output=True,
            text=True,
            timeout=80,
        )
        return {"exit_code": result.returncode, "output_tail": (result.stdout + result.stderr)[-16000:]}

    with httpx.Client(timeout=httpx.Timeout(180, connect=10)) as client:
        for call in range(20):
            response = client.post(
                spec.base_url.rstrip("/") + "/chat/completions",
                headers={"Authorization": "Bearer " + spec.api_key()},
                json={
                    "model": spec.upstream_model,
                    "messages": messages,
                    "tools": tool_catalog,
                    "max_tokens": 4096,
                    "thinking": {"type": "disabled"},
                },
            )
            if response.status_code != 200:
                raise RuntimeError("model HTTP status " + str(response.status_code))
            body = response.json()
            for key in usage:
                usage[key] += body.get("usage", {}).get(key, 0)
            message = body["choices"][0]["message"]
            messages.append(message)
            calls = message.get("tool_calls") or []
            print(json.dumps({"call": call + 1, "tools": [t["function"]["name"] for t in calls]}), flush=True)
            if not calls:
                break
            for tool in calls:
                name = tool["function"]["name"]
                values = json.loads(tool["function"]["arguments"])
                if name == "search_capabilities":
                    skills = [{"name": name, "description": item.get("description", "")} for name, item in prompts.skills().items()]
                    result = discovery.search(values["query"], values.get("limit", 5), skills)
                elif name == "load_skill":
                    skill = values["name"]
                    try:
                        result = internal_load(skill) if skill.startswith("platform.") else prompts.load_skill(skill)
                    except (KeyError, ValueError):
                        result = {"error": "skill unavailable; use search_capabilities to discover actual skills"}
                elif name == "exec":
                    command = values["cmd"]
                    commands.append(command)
                    result = docker("sh", "-lc", command)
                elif name == "read_file":
                    result = docker(
                        "python", "-c", "import pathlib,sys; print(pathlib.Path(sys.argv[1]).read_text()[:16000])", values["path"]
                    )
                elif name == "write_file":
                    result = docker(
                        "python",
                        "-c",
                        "import pathlib,sys; p=pathlib.Path(sys.argv[1]); "
                        "p.parent.mkdir(parents=True,exist_ok=True); p.write_text(sys.argv[2])",
                        values["path"],
                        values["content"],
                    )
                elif name == "deliver_file":
                    result = {
                        "isolated_verification": True,
                        "file_path": values["path"],
                        "delivered": False,
                        "note": "验证不写业务库，文件保存在容器outputs；无需重复交付，可结束回答。",
                    }
                elif name == "update_plan":
                    result = {"ok": True}
                else:
                    result = {"error": "tool not active or unsupported in isolated verification"}
                print(json.dumps({"tool": name, "args": values, "result": result}, ensure_ascii=False), flush=True)
                messages.append({"role": "tool", "tool_call_id": tool["id"], "content": json.dumps(result, ensure_ascii=False)})
        else:
            raise RuntimeError("20-call cap reached")
    check = docker(
        "python",
        "-c",
        """
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
import json
text='\\n'.join(p.text for p in Document('/workspace/outputs/report.docx').paragraphs)
assert '175' in text and '120' not in text and 'KEEP-UNCHANGED' in text, text
w=load_workbook('/workspace/outputs/budget.xlsx'); s=w.active
assert s['B2'].value == 175 and s['B3'].value == 80 and s['B4'].value == '=SUM(B2:B3)'
assert len(list(Path('/workspace').rglob('manifest.json'))) >= 4, 'missing before/after reader results'
print('MODEL ARTIFACT CHECK PASS')
""",
    )
    assert check["exit_code"] == 0, check
    assert any("office-read" in command for command in commands), "model did not use office-read"
    print(json.dumps({"model": spec.upstream_model, "calls": call + 1, "usage": usage, "artifact_check": "passed"}))


if __name__ == "__main__":
    main()
