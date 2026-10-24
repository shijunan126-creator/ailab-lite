"""Lightweight template sketch, only in a limited parser child."""

import io


def _sketch(prs, output):
    from PIL import Image, ImageDraw, ImageFont

    width, height = 480, max(1, round(480 * prs.slide_height / prs.slide_width))
    height = min(height, 960)
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    slide = prs.slides[0] if len(prs.slides) else prs.slide_layouts[0]
    sx, sy = width / prs.slide_width, height / prs.slide_height
    Image.MAX_IMAGE_PIXELS = 16_000_000
    for shape in list(slide.shapes)[:200]:
        x, y = round(shape.left * sx), round(shape.top * sy)
        w, h = max(1, round(shape.width * sx)), max(1, round(shape.height * sy))
        if w * h > 4 * width * height:
            continue
        if shape.shape_type == 13:  # embedded picture; decoder stays within child limits
            with Image.open(io.BytesIO(shape.image.blob)) as image:
                image.thumbnail((min(w, width), min(h, height)))
                canvas.paste(image.convert("RGB"), (x, y))
        elif shape.has_text_frame:
            content = shape.text[:1000]
            if not content:
                continue
            size = max(8, min(32, round((shape.text_frame.paragraphs[0].font.size or 180000) * sy)))
            try:
                font = ImageFont.truetype("DejaVuSans.ttf", size)
            except OSError:
                font = ImageFont.load_default(size=size)
            draw.multiline_text((x, y), content, fill="#27364a", font=font, spacing=2)
        else:
            try:
                color = shape.fill.fore_color.rgb
                if color:
                    draw.rectangle((x, y, x + w, y + h), fill="#" + str(color))
            except (AttributeError, TypeError):
                pass
    canvas.save(output / "thumb.jpg", "JPEG", quality=82, optimize=True)


def sketch(source, output):
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.package import Package
    from pptx.parts.presentation import PresentationPart

    PartFactory.part_type_for[CT.PML_TEMPLATE_MAIN] = PresentationPart
    prs = Package.open(source).main_document_part.presentation
    _sketch(prs, output)
    return {"thumbnail": "sketch"}
