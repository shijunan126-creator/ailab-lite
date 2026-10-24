"""Only imported inside the limited child process (never in a Worker thread)."""

from __future__ import annotations

import warnings
from pathlib import Path, PurePath

MAX_TEXT = 200_000
TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".log",
    ".py",
    ".sh",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".css",
    ".sql",
    ".go",
    ".rs",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".toml",
    ".ini",
}


def images(source: Path, out: Path, options: dict) -> dict:
    from PIL import Image, ImageOps

    if source.stat().st_size > options.get("max_bytes", 20 * 1024 * 1024):
        raise ValueError("input_limit")
    Image.MAX_IMAGE_PIXELS = options.get("pillow_pixels", 128_000_000)
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(source) as image:
            if image.format in {"EPS", "WMF", "PDF"}:
                raise ValueError("corrupt")
            bound = options.get("jpeg_pixels", 128_000_000) if image.format == "JPEG" else options.get("decoded_pixels", 16_000_000)
            if image.width * image.height > bound:
                raise ValueError("pixel_limit")
            image.seek(0)
            image.draft("RGB", (2048, 2048))
            image.thumbnail((2048, 2048))
            with ImageOps.exif_transpose(image) as oriented:
                mode = "RGBA" if "A" in oriented.getbands() or "transparency" in oriented.info else "RGB"
                with oriented.convert(mode) as display:
                    display.save(out / "display.webp", "WEBP", quality=80, method=4)
                    display.thumbnail((480, 480))
                    display.save(out / "thumb.webp", "WEBP", quality=75, method=4)
    return {"image": "display.webp", "thumbnail": "thumb.webp"}


def _pdf(data: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(data)
    pages = len(reader.pages)
    text = "".join(f"\f[page:{index + 1}]\n{page.extract_text() or ''}" for index, page in enumerate(reader.pages[:20]))
    return {"kind": "pdf", "pages": pages, "chars": len(text), "truncated": len(text) > MAX_TEXT, "text": text[:MAX_TEXT]}


def _docx(data: Path) -> dict:
    import docx

    doc = docx.Document(data)
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + "\t".join(cell.text for cell in row.cells)
    return {
        "kind": "docx",
        "paragraphs": len(doc.paragraphs),
        "chars": len(text),
        "truncated": len(text) > MAX_TEXT,
        "text": text[:MAX_TEXT],
    }


def _pptx(data: Path) -> dict:
    from pptx import Presentation

    prs = Presentation(data)
    chunks: list[str] = []
    for index, slide in enumerate(prs.slides):
        chunks.append(f"\f[page:{index + 1}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                chunks.append(shape.text)
    text = "\n".join(chunks)
    return {"kind": "pptx", "slides": len(prs.slides), "chars": len(text), "truncated": len(text) > MAX_TEXT, "text": text[:MAX_TEXT]}


def _xlsx(data: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(data, read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets[:20]:
        chunks.append(f"# {ws.title}")
        for row in ws.iter_rows(max_row=200, max_col=40, values_only=True):
            chunks.append("\t".join("" if v is None else str(v) for v in row))
    text = "\n".join(chunks)
    wb.close()
    return {"kind": "xlsx", "sheets": len(wb.worksheets), "chars": len(text), "truncated": len(text) > MAX_TEXT, "text": text[:MAX_TEXT]}


def extract_path(data: Path, filename: str, content_type: str) -> dict:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf" or content_type == "application/pdf":
        return _pdf(data)
    if suffix == ".docx":
        return _docx(data)
    if suffix == ".pptx":
        return _pptx(data)
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx(data)

    if suffix in TEXT_SUFFIXES or content_type.startswith("text/"):
        with data.open("rb") as stream:
            prefix = stream.read(4)
        encodings = (
            ["utf-32"]
            if prefix.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff"))
            else ["utf-16"]
            if prefix.startswith((b"\xff\xfe", b"\xfe\xff"))
            else ["utf-8-sig", "gb18030"]
        )
        for encoding in [*encodings, "utf-8"]:
            try:
                with data.open(encoding=encoding, errors="replace" if encoding == "utf-8" else "strict") as stream:
                    text = stream.read(MAX_TEXT + 1)
                break
            except UnicodeDecodeError:
                continue
        return {"kind": "text", "chars": len(text), "truncated": len(text) > MAX_TEXT, "text": text[:MAX_TEXT]}

    return {"kind": "binary", "chars": 0, "text": "", "bytes": data.stat().st_size}


def run(operation: str, source: Path, output: Path, options: dict) -> dict:
    filename = options.get("filename", source.name)
    from ..sandbox_tools.render_preview import check_archive

    check_archive(source)
    if operation == "template":
        return analyze_template(source, output)
    if operation == "brand":
        return analyze_brand(source, output, options["kind"])
    if operation == "ocr_pdf":
        from .ocr import low_text_pages

        return low_text_pages(source, output, **options)
    if operation == "images":
        return images(source, output, options)
    if operation == "sheets":
        from ..sandbox_tools.render_preview import sheets

        return sheets(source, output, recalculate=False)
    if operation == "extract":
        from .. import skillfiles

        info = extract_path(source, filename, options.get("content_type", ""))
        try:
            found = skillfiles.inspect(source.read_bytes(), filename) if Path(filename).suffix.lower() in {".zip", ".md"} else None
            if found:
                info["skill"] = found
        except skillfiles.SkillFileError as exc:
            info["skill_error"] = str(exc)
        return info
    raise ValueError("unknown operation")


def analyze_brand(source: Path, output: Path, kind: str, presentation=None) -> dict:
    if kind == "logo":
        from PIL import Image, ImageOps

        Image.MAX_IMAGE_PIXELS = 16_000_000
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(source) as image:
                if image.format not in {"PNG", "JPEG"}:
                    raise ValueError("corrupt")
                image.load()
                with ImageOps.exif_transpose(image) as oriented:
                    oriented.convert("RGBA").save(output / "logo.png", "PNG")
                size = (output / "logo.png").stat().st_size
                if size > 2 * 1024 * 1024:
                    raise ValueError("output_limit")
                return {"width": image.width, "height": image.height, "png_bytes": size}
    if kind == "docx":
        from docx import Document

        doc = Document(source)
        return {"styles": [style.name[:120] for style in list(doc.styles)[:50]]}
    if kind == "pptx":
        import zipfile

        from lxml import etree
        from pptx import Presentation

        prs = presentation if presentation is not None else Presentation(source)
        themes = []
        with zipfile.ZipFile(source) as archive:
            for name in sorted(n for n in archive.namelist() if n.startswith("ppt/theme/theme") and n.endswith(".xml"))[:20]:
                theme = etree.fromstring(archive.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                colors = {}
                for entry in theme.findall(".//a:clrScheme/*", ns):
                    if len(entry):
                        colors[etree.QName(entry).localname] = entry[0].get("val") or entry[0].get("lastClr")
                fonts = {
                    key: theme.xpath(f"string(.//a:{key}/a:latin/@typeface)", namespaces=ns)[:120] for key in ("majorFont", "minorFont")
                }
                themes.append({"colors": colors, "fonts": fonts})
        return {
            "slide_width": prs.slide_width,
            "slide_height": prs.slide_height,
            "slides": len(prs.slides),
            "layout_details": [
                {
                    "name": layout.name[:120],
                    "placeholders": [
                        {"idx": p.placeholder_format.idx, "type": int(p.placeholder_format.type)} for p in layout.placeholders
                    ],
                }
                for master in prs.slide_masters
                for layout in master.slide_layouts
            ][:100],
            "layouts": [layout.name[:120] for master in prs.slide_masters for layout in master.slide_layouts][:100],
            "master_count": len(prs.slide_masters),
            "themes": themes,
        }
    raise ValueError("unknown brand kind")


def analyze_template(source: Path, output: Path) -> dict:
    import zipfile

    from .template_preview import thumbnail

    with zipfile.ZipFile(source) as archive:
        types = archive.read("[Content_Types].xml")
        if b"macroenabled" in types.lower() or any("vbaproject" in n.lower() for n in archive.namelist()):
            raise ValueError("macro_template")
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.package import Package
    from pptx.parts.presentation import PresentationPart

    # POTX uses the same presentation XML with a distinct package main content type.
    # Open it directly: rewriting a 30 MiB package would exceed the parser output cap.
    PartFactory.part_type_for[CT.PML_TEMPLATE_MAIN] = PresentationPart
    part = Package.open(source).main_document_part
    if part.content_type not in {CT.PML_TEMPLATE_MAIN, CT.PML_PRESENTATION_MAIN}:
        raise ValueError("unsupported_template")
    prs = part.presentation
    result = analyze_brand(source, output, "pptx", presentation=prs)
    thumbnail(prs, output)
    return result
