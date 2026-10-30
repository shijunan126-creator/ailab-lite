"""Template upload→outbox→Worker→API→Agent→Broker, then clean restored workspace."""

import io
import time

import httpx
import pytest
from pptx import Presentation

from .conftest import PASSWORD, USERNAME
from .test_portal import _pick_skill

pytestmark = pytest.mark.e2e


def wait(client, path, predicate, seconds=60):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        response = client.get(path)
        assert response.status_code == 200, response.text
        value = response.json()
        if predicate(value):
            return value
        time.sleep(0.2)
    raise AssertionError(f"condition not reached: {path}; last={value}")


def test_template_upload_run_and_next_turn_reset(stack):
    with httpx.Client(base_url=stack.api, timeout=15) as client:
        login = client.post("/v1/auth/login", json={"username": USERNAME, "password": PASSWORD})
        assert login.status_code == 200, login.text
        client.headers["Authorization"] = "Bearer " + login.json()["token"]
        stream = io.BytesIO()
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Template sample"
        prs.save(stream)
        response = client.post("/v1/templates", files={"upload": ("company.pptx", stream.getvalue())}, data={"scope": "company"})
        assert response.status_code == 201, response.text
        template_id = response.json()["id"]
        try:
            tpl = wait(client, f"/v1/templates/{template_id}", lambda row: row["status"] != "pending")
            assert tpl["status"] == "ready", tpl
            thumb = client.get(tpl["thumbnail_url"])
            assert thumb.status_code == 200 and "immutable" in thumb.headers["cache-control"]
            conv = client.post("/v1/conversations", json={}).json()
            path = f"/v1/conversations/{conv['id']}/messages"
            first = client.post(path, json={"prompt": "PPT模板链路-有模板", "template": {"id": template_id}})
            assert first.status_code == 201, first.text
            run = wait(client, "/v1/runs/" + first.json()["id"], lambda r: r["status"] in {"succeeded", "failed"})
            assert run["status"] == "succeeded" and "TEMPLATE_PRESENT" in run["summary"], run
            assert run["template"] == {"source": "template", "name": "company"}
            second = client.post(path, json={"prompt": "PPT模板链路-无模板"})
            assert second.status_code == 201, second.text
            plain = wait(client, "/v1/runs/" + second.json()["id"], lambda r: r["status"] in {"succeeded", "failed"})
            assert plain["status"] == "succeeded" and "TEMPLATE_ABSENT" in plain["summary"], plain
            assert plain["template"] is None
            client.delete("/v1/conversations/" + conv["id"])
        finally:
            assert client.delete(f"/v1/templates/{template_id}").status_code == 204


@pytest.mark.parametrize("scope", ["company", "personal"])
@pytest.mark.parametrize("preview_result", ["ready", "failed"])
def test_template_preview_updates_without_reopening(page, stack, scope, preview_result):
    """A usable template keeps its sketch until the independent preview settles."""
    from playwright.sync_api import expect

    from .conftest import login

    template = {
        "id": "tpl_preview_regression",
        "kind": "pptx",
        "scope": scope,
        "name": "Delayed preview",
        "status": "pending",
        "preview_status": "unsupported",
        "reason": None,
        "thumbnail_url": None,
        "layouts": 2,
        "slides": 1,
        "size": 1024,
        "filename": "delayed.pptx",
        "created_at": "2026-10-24T00:00:00Z",
        "can_edit": True,
    }
    calls = []
    sketch = "/v1/templates/tpl_preview_regression/thumbnail"
    rendered = sketch + "?v=rendered"

    def templates(route):
        if route.request.url.split("?")[0].endswith("/tpl_preview_regression"):
            calls.append(route.request.url)
            # A transient read failure must not permanently stop lifecycle tracking.
            if len(calls) == 1:
                route.fulfill(status=503, json={"detail": "temporarily unavailable"})
                return
            template.update(status="ready", preview_status="pending", thumbnail_url=sketch)
            if len(calls) >= 3:
                template.update(preview_status=preview_result, thumbnail_url=rendered if preview_result == "ready" else sketch)
            route.fulfill(json=template)
        else:
            route.fulfill(json={"company": [template] if scope == "company" else [], "personal": [template] if scope == "personal" else []})

    page.route("**/api/v1/templates*", templates)
    page.route("**/api/v1/templates/tpl_preview_regression", templates)
    login(page, stack)
    if scope == "company":
        page.goto(stack.portal + "/admin?t=brand")
    else:
        _pick_skill(page, "PPT 设计")
        page.get_by_role("button", name="模板", exact=True).click()
    container = (
        page.locator(".brand-slot").filter(has_text="Delayed preview")
        if scope == "company"
        else page.locator(".tpl-item").filter(has_text="Delayed preview")
    )
    expect(container).to_be_visible()
    image = container.locator("img")
    expect(image).to_have_attribute("src", "/api" + sketch, timeout=15000)
    # The template is selectable while preview_status remains pending.
    if scope == "personal":
        expect(container.locator(".tpl-card")).to_be_enabled()
    with page.expect_response(
        lambda response: (
            response.url.endswith("/tpl_preview_regression")
            and response.status == 200
            and response.json().get("preview_status") == preview_result
        ),
        timeout=15000,
    ) as terminal:
        expect(image).to_have_attribute("src", "/api" + sketch)
    terminal.value.finished()
    expect(image).to_have_attribute("src", "/api" + (rendered if preview_result == "ready" else sketch))
    page.clock.install()
    page.clock.run_for(10000)
    assert len(calls) == 3, "terminal previews must stop polling"
