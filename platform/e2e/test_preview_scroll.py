"""Both preview surfaces scroll native document pixels, with no wheel page flipping."""

import io
import subprocess
import sys
import uuid

import pytest
from playwright.sync_api import expect

from .conftest import select_model, wait_models

pytestmark = pytest.mark.e2e


def test_library_pages_scroll_both_ways_and_contain_overscroll(logged_in, stack):
    from docx import Document

    doc = Document()
    for number in range(4):
        if number:
            doc.add_page_break()
        doc.add_heading(f"Page {number + 1}", 0)
        doc.add_paragraph("Scroll acceptance")
    source = io.BytesIO()
    doc.save(source)
    name = f"scroll-{uuid.uuid4().hex[:6]}.docx"
    page = logged_in
    response = page.request.post(
        stack.portal + "/api/v1/files",
        multipart={
            "upload": {
                "name": name,
                "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "buffer": source.getvalue(),
            }
        },
    )
    assert response.ok
    fid = response.json()["id"]
    page.goto(stack.portal + "/files")
    row = page.locator(".lib-row", has_text=name)
    row.get_by_role("link", name=f"下载 {name}", exact=True).wait_for(timeout=30000)
    # Deterministic rendered pages isolate native scrolling from host Office/font availability.
    script = """
import hashlib,io,sys
from PIL import Image,ImageDraw
from ailab_platform.db import SessionLocal
from ailab_platform.models import FileObject
from ailab_platform.storage import get_storage
fid=sys.argv[1];storage=get_storage();names=[];assets={}
for i in range(1,5):
 image=Image.new('RGB',(800,1100),'white');ImageDraw.Draw(image).text((50,50),f'Page {i}',fill='black')
 body=io.BytesIO();image.save(body,format='JPEG');name=f'p-{i:02d}.jpg';names.append(name)
 data=body.getvalue();key=f'derived/{fid}/preview/{name}'
 storage.put_bytes(key,data,'image/jpeg')
 assets[name]={'key':key,'sha256':hashlib.sha256(data).hexdigest(),'content_type':'image/jpeg'}
with SessionLocal() as db:
 row=db.get(FileObject,fid)
 meta=dict(row.meta or {})
 meta['preview']={'status':'ready','kind':'pages','pages':names,'page_count':4}
 meta['preview_assets']=assets
 row.meta=meta
 db.commit()
"""
    subprocess.run([sys.executable, "-c", script, fid], env=stack.env, check=True, timeout=15)
    row.get_by_role("button", name=name, exact=True).click()
    viewer = page.get_by_role("dialog", name=name)
    scroll = viewer.locator(".document-scroll")
    expect(scroll.locator("img")).to_have_count(4, timeout=90000)
    first = scroll.locator("img").first
    page.wait_for_function("img => img.complete && img.naturalWidth > 0", arg=first.element_handle())
    outer = page.evaluate("window.scrollY")
    scroll.hover()
    page.mouse.wheel(0, 100)
    page.wait_for_function("el => el.scrollTop > 0", arg=scroll.element_handle())
    assert scroll.evaluate("el => el.scrollTop") < 500  # a small wheel gesture cannot jump a whole page
    page.mouse.wheel(0, 1600)
    expect(viewer.locator(".document-counter")).not_to_have_text("1 / 4")
    page.mouse.wheel(0, -10000)
    page.wait_for_function("el => el.scrollTop === 0", arg=scroll.element_handle())
    expect(viewer.locator(".document-counter")).to_have_text("1 / 4")
    assert page.evaluate("window.scrollY") == outer
    assert scroll.evaluate("el => getComputedStyle(el).overscrollBehaviorY") == "contain"
    page.keyboard.press("Escape")


def test_chat_preview_uses_the_same_continuous_pages(logged_in, stack):
    page = logged_in
    wait_models(page)
    select_model(page, "test-fast")
    page.fill("textarea", "你好")
    page.keyboard.press("Enter")
    page.locator(".turn-meta").wait_for(timeout=30000)
    cid = page.url.rsplit("/", 1)[-1]
    rid = page.request.get(stack.portal + "/api/v1/conversations/" + cid).json()["turns"][-1]["id"]
    # Deterministic rendered pages isolate native scrolling from host Office/font availability.
    script = """
import io,sys
from PIL import Image,ImageDraw
from ailab_platform.db import SessionLocal
from ailab_platform.models import Run
from ailab_platform.storage import get_storage
rid=sys.argv[1];storage=get_storage();names=[]
for i in range(1,4):
 image=Image.new('RGB',(800,1100),'white');ImageDraw.Draw(image).text((50,50),f'Page {i}',fill='black')
 body=io.BytesIO();image.save(body,format='JPEG');name=f'1/p-{i}.jpg';names.append(name)
 storage.put_bytes(f'previews/{rid}/{name}',body.getvalue(),'image/jpeg')
with SessionLocal() as db:
 run=db.get(Run,rid);run.result={**run.result,'previews':[{'filename':'scroll.docx','kind':'pages','version':1,'pages':names}]};db.commit()
"""
    subprocess.run([sys.executable, "-c", script, rid], env=stack.env, check=True, timeout=15)
    page.reload()
    page.locator(".composer-dock .preview-reopen").click()
    pane = page.get_by_role("complementary", name="文件预览")
    scroll = pane.locator(".document-scroll")
    expect(scroll.locator("img")).to_have_count(3)
    first = scroll.locator("img").first
    page.wait_for_function("img => img.complete && img.naturalWidth > 0", arg=first.element_handle())
    before = page.locator(".messages").evaluate("el => el.scrollTop")
    scroll.hover()
    page.mouse.wheel(0, 1200)
    page.wait_for_function("el => el.scrollTop > 0", arg=scroll.element_handle())
    expect(pane.locator(".document-counter")).not_to_have_text("1 / 3")
    page.mouse.wheel(0, -10000)
    page.wait_for_function("el => el.scrollTop === 0", arg=scroll.element_handle())
    expect(pane.locator(".document-counter")).to_have_text("1 / 3")
    assert page.locator(".messages").evaluate("el => el.scrollTop") == before
