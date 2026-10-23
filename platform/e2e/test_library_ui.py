"""Browser tests of the library viewer, the trash, version history and the admin console.

Screenshots of each state go to AILAB_E2E_SHOTS (when set) for a visual review.
"""

from __future__ import annotations

import io
import os
import shutil
import uuid
from pathlib import Path

import pytest

pytestmark = pytest.mark.e2e

SHOTS = Path(os.environ["AILAB_E2E_SHOTS"]) if os.environ.get("AILAB_E2E_SHOTS") else None


def _shot(page, name: str) -> None:
    if SHOTS:
        SHOTS.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(SHOTS / f"{name}.png"))


def _upload(page, stack, tmp_path, name: str, data: bytes) -> None:
    path = tmp_path / name
    path.write_bytes(data)
    page.goto(stack.portal + "/files")
    page.set_input_files("input[type=file]", str(path))
    # usable once the Worker has scanned it: the download link appears
    page.locator(".lib-row", has_text=name).get_by_role("link", name=f"下载 {name}", exact=True).wait_for(timeout=30000)


def _open(page, name: str):
    page.locator(".lib-row", has_text=name).get_by_role("button", name=name, exact=True).click()
    viewer = page.get_by_role("dialog", name=name)
    viewer.wait_for()
    return viewer


def test_markdown_opens_in_place_and_escape_closes(logged_in, stack, tmp_path) -> None:
    page = logged_in
    name = f"notes-{uuid.uuid4().hex[:6]}.md"
    _upload(page, stack, tmp_path, name, "# 季度计划\n\n- 第一项\n- 第二项\n".encode())
    viewer = _open(page, name)
    viewer.get_by_role("heading", name="季度计划").wait_for(timeout=15000)
    assert viewer.locator(".md li").count() == 2
    _shot(page, "viewer-markdown")
    page.keyboard.press("Escape")
    viewer.wait_for(state="detached")


def test_a_workbook_previews_as_sheets(logged_in, stack, tmp_path) -> None:
    import openpyxl

    page = logged_in
    book = openpyxl.Workbook()
    book.active.title = "销售"
    book.active.append(["地区", "金额"])
    book.active.append(["华东", 1200])
    data = io.BytesIO()
    book.save(data)
    name = f"sales-{uuid.uuid4().hex[:6]}.xlsx"
    _upload(page, stack, tmp_path, name, data.getvalue())
    viewer = _open(page, name)
    viewer.locator(".viewer-sheet td", has_text="华东").wait_for(timeout=30000)  # rendered by the Worker on first open
    assert viewer.get_by_role("button", name="销售").is_visible()
    _shot(page, "viewer-sheets")
    page.keyboard.press("Escape")


@pytest.mark.skipif(not all(shutil.which(x) for x in ("soffice", "pdftoppm", "pdfinfo", "timeout")), reason="Office tools unavailable")
def test_a_document_previews_as_pages_rendered_in_a_sandbox(logged_in, stack, tmp_path) -> None:
    from docx import Document

    page = logged_in
    doc = Document()
    doc.add_heading("项目周报", 0)
    doc.add_paragraph("本周完成了资料库预览。")
    data = io.BytesIO()
    doc.save(data)
    name = f"weekly-{uuid.uuid4().hex[:6]}.docx"
    _upload(page, stack, tmp_path, name, data.getvalue())
    viewer = _open(page, name)
    image = viewer.locator(".viewer-pages img").first
    image.wait_for(timeout=90000)
    page.wait_for_function("(img) => img.complete && img.naturalWidth > 0", arg=image.element_handle(), timeout=30000)
    _shot(page, "viewer-pages")
    page.keyboard.press("Escape")


def test_arrows_step_through_the_list(logged_in, stack, tmp_path) -> None:
    page = logged_in
    first, second = (f"{tag}-{uuid.uuid4().hex[:6]}.txt" for tag in ("alpha", "beta"))
    _upload(page, stack, tmp_path, first, b"first")
    _upload(page, stack, tmp_path, second, b"second")
    _open(page, second)  # newest first: the older file is next
    page.keyboard.press("ArrowRight")
    page.get_by_role("dialog", name=first).wait_for()
    page.keyboard.press("ArrowLeft")
    page.get_by_role("dialog", name=second).wait_for()
    page.keyboard.press("Escape")


def test_delete_offers_undo_and_the_trash_restores_and_purges(logged_in, stack, tmp_path) -> None:
    page = logged_in
    name = f"draft-{uuid.uuid4().hex[:6]}.txt"
    _upload(page, stack, tmp_path, name, b"draft")
    row = page.locator(".lib-row", has_text=name)
    row.get_by_role("button", name=f"更多操作 {name}", exact=True).click()
    row.get_by_role("menuitem", name="删除", exact=True).click()  # no confirmation: it only goes to the trash
    toast = page.get_by_role("status").filter(has_text="已移到回收站")
    toast.wait_for()
    _shot(page, "trash-undo")
    toast.get_by_role("button", name="撤销").click()
    row.wait_for()

    row.get_by_role("button", name=f"更多操作 {name}", exact=True).click()
    row.get_by_role("menuitem", name="删除", exact=True).click()
    row.wait_for(state="detached")
    page.get_by_role("tab", name="回收站").click()
    trashed = page.locator(".lib-row", has_text=name)
    trashed.wait_for()
    assert "天后永久删除" in trashed.inner_text()
    _shot(page, "trash-list")
    trashed.get_by_role("button", name=f"恢复 {name}").click()
    trashed.wait_for(state="detached")
    page.get_by_role("tab", name="全部").click()
    page.locator(".lib-row", has_text=name).wait_for()

    page.locator(".lib-row", has_text=name).get_by_role("button", name=f"更多操作 {name}", exact=True).click()
    page.locator(".lib-row", has_text=name).get_by_role("menuitem", name="删除", exact=True).click()
    page.get_by_role("tab", name="回收站").click()
    page.on("dialog", lambda d: d.accept())
    page.locator(".lib-row", has_text=name).get_by_role("button", name=f"永久删除 {name}").click()
    page.locator(".lib-row", has_text=name).wait_for(state="detached")
    files = page.request.get(f"{stack.portal}/api/v1/files?trashed=true").json()
    assert all(f["filename"] != name for f in files)


def test_an_old_version_can_be_viewed_and_restored(logged_in, stack, tmp_path) -> None:
    page = logged_in
    name = f"plan-{uuid.uuid4().hex[:6]}.md"
    _upload(page, stack, tmp_path, name, "# 第一版\n".encode())
    first = next(f for f in page.request.get(f"{stack.portal}/api/v1/files?q={name}").json() if f["filename"] == name)
    upload = {"upload": {"name": name, "mimeType": "text/markdown", "buffer": "# 第二版\n".encode()}}
    second = page.request.post(f"{stack.portal}/api/v1/files?parent_id={first['id']}", multipart=upload).json()
    page.reload()
    page.locator(".lib-row", has_text=name).get_by_role("link", name=f"下载 {name}", exact=True).wait_for(timeout=30000)
    viewer = _open(page, name)
    viewer.get_by_role("heading", name="第二版").wait_for(timeout=15000)
    viewer.get_by_role("button", name="版本历史").click()
    viewer.locator(".viewer-menu-item", has_text="v1").click()
    viewer.get_by_role("heading", name="第一版").wait_for(timeout=15000)
    _shot(page, "viewer-old-version")
    viewer.get_by_role("button", name="恢复此版本").click()
    page.get_by_role("dialog", name=name).get_by_role("button", name="版本历史").filter(has_text="v3").wait_for(timeout=15000)
    versions = page.request.get(f"{stack.portal}/api/v1/files/{second['id']}/versions").json()
    assert [v["version"] for v in versions] == [3, 2, 1]


def test_the_admin_console_manages_departments_and_shows_the_audit(logged_in, stack) -> None:
    page = logged_in
    page.goto(stack.portal + "/")
    page.locator(".user-chip").click()
    page.get_by_role("menuitem", name="管理").click()
    page.wait_for_url("**/admin")
    page.locator(".user-row").first.wait_for()
    _shot(page, "admin-users")

    dept = f"研发部-{uuid.uuid4().hex[:4]}"
    page.get_by_role("tab", name="部门").click()
    page.get_by_role("button", name="新建部门").click()
    page.get_by_role("dialog", name="新建部门").get_by_label("名称").fill(dept)
    page.get_by_role("dialog", name="新建部门").get_by_role("button", name="创建").click()
    page.locator(".admin-row", has_text=dept).wait_for()

    page.get_by_role("tab", name="额度").click()
    page.get_by_role("button", name="新建额度").click()
    dialog = page.get_by_role("dialog", name="新建额度")
    dialog.get_by_label("名称").fill(f"标准-{uuid.uuid4().hex[:4]}")
    dialog.get_by_role("button", name="创建").click()
    page.locator(".policy-card").first.wait_for()
    _shot(page, "admin-policies")

    page.get_by_role("tab", name="审计").click()
    page.locator(".audit-item", has_text="新建部门").first.wait_for()
    _shot(page, "admin-audit")

    page.get_by_role("tab", name="用量").click()
    page.get_by_role("radio", name="本月").wait_for()
    _shot(page, "admin-usage")
