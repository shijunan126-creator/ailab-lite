"""Office skill checkers catch concrete defects before delivery."""

import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[2] / "images" / "sbx-office" / "skills"


def run(kind, source):
    return subprocess.run(
        [sys.executable, str(ROOT / f"platform.{kind}" / "scripts" / "check.py"), str(source)], capture_output=True, text=True, timeout=120
    )


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice unavailable")
def test_xlsx_checker_recalculates_and_reports_formula_error(tmp_path):
    workbook = Workbook()
    workbook.active["A1"] = "=1/0"
    path = tmp_path / "bad.xlsx"
    workbook.save(path)
    result = run("xlsx", path)
    assert result.returncode == 1
    assert "A1" in result.stdout and "#DIV/0!" in result.stdout


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice unavailable")
def test_xlsx_checker_detects_empty_chart_data(tmp_path):
    from openpyxl.chart import BarChart, Reference

    workbook = Workbook()
    chart = BarChart()
    chart.add_data(Reference(workbook.active, min_col=1, min_row=1, max_row=3))
    workbook.active.add_chart(chart, "C1")
    path = tmp_path / "empty-chart.xlsx"
    workbook.save(path)
    result = run("xlsx", path)
    assert result.returncode == 1
    assert "empty or broken" in result.stdout


def test_pptx_checker_reports_overflow_and_off_slide_shape(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.2))
    shape.text_frame.paragraphs[0].text = "This deliberately long text cannot fit into this tiny shape. " * 4
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    slide.shapes.add_textbox(deck.slide_width, 0, Inches(1), Inches(1)).text = "outside"
    path = tmp_path / "bad.pptx"
    deck.save(path)
    result = run("pptx", path)
    assert result.returncode == 0
    assert "text overflow" in result.stdout and "off-slide" in result.stdout


def test_docx_and_pdf_checks_report_usable_documents(tmp_path):
    doc = Document()
    doc.add_paragraph("Content")
    word = tmp_path / "doc.docx"
    doc.save(word)
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    pdf = tmp_path / "doc.pdf"
    with pdf.open("wb") as stream:
        writer.write(stream)
    assert run("docx", word).returncode == 0
    result = run("pdf", pdf)
    assert result.returncode == 0 and "1 pages" in result.stdout
