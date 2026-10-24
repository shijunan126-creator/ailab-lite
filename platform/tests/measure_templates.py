"""Small local template parser probe; no developer profile, network or model."""

import json
import statistics
import tempfile
import time
from pathlib import Path

from pptx import Presentation

from ailab_platform.worker.isolate import execute


def main():
    with tempfile.TemporaryDirectory(prefix="ailab-template-measure-") as tmp:
        root = Path(tmp)
        source = root / "template.pptx"
        prs = Presentation()
        for index in range(20):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"Template sample {index}"
            slide.placeholders[1].text = "Example body\n" * 8
        prs.save(source)
        samples = []
        for index in range(5):
            output = root / str(index)
            started = time.perf_counter()
            result = execute("template", source, output, measure=True)
            samples.append(
                {
                    **result["measurement"],
                    "wall_seconds": time.perf_counter() - started,
                    "output_bytes": sum(p.stat().st_size for p in output.iterdir() if p.is_file()),
                }
            )
        print(
            json.dumps(
                {
                    "slides": 20,
                    "input_bytes": source.stat().st_size,
                    "cpu_seconds_median": statistics.median(r["cpu_seconds"] for r in samples),
                    "wall_seconds_median": statistics.median(r["wall_seconds"] for r in samples),
                    "peak_rss_bytes_max": max(r["peak_rss_bytes"] for r in samples),
                    "output_bytes_max": max(r["output_bytes"] for r in samples),
                },
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
