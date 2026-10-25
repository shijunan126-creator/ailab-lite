"""PPT template registry, frozen run grants, queue and sandbox isolation."""

import importlib.util
import io
import json
import shutil
import uuid
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select, update

from ailab_platform import accounts
from ailab_platform.api.routes_internal import _bound_file
from ailab_platform.config import settings
from ailab_platform.db import SessionLocal
from ailab_platform.models import FileObject, Run, Template, utcnow
from ailab_platform.quota import library_bytes
from ailab_platform.service_auth import ServicePrincipal
from ailab_platform.worker.pipeline import process_file
from tests.run_claims import attempt_headers
from tests.test_broker_lease import _ApiStub, _manager_with
from tests.test_projects_brand import SERVICE, WorkerAPI, start
from tests.test_projects_brand import api as api


@pytest.fixture(autouse=True)
def isolate_registry(api):
    with SessionLocal() as db:
        db.execute(update(Template).values(deleted_at=utcnow()))
        db.commit()


def pptx(slides=1, potx=False, macro=False):
    out = io.BytesIO()
    prs = Presentation()
    for _ in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "SAMPLE TEXT TO REMOVE"
    prs.save(out)
    if not potx and not macro:
        return out.getvalue()
    converted = io.BytesIO()
    with zipfile.ZipFile(out) as src, zipfile.ZipFile(converted, "w") as dst:
        for entry in src.infolist():
            value = src.read(entry)
            if entry.filename == "[Content_Types].xml":
                value = value.replace(
                    b"presentationml.presentation.main+xml",
                    b"presentationml.template.main+xml" if potx else b"presentation.macroEnabled.main+xml",
                )
            dst.writestr(entry, value)
    return converted.getvalue()


def upload(api, *, scope="personal", name="demo.pptx", data=None):
    response = api.post("/v1/templates", files={"upload": (name, pptx() if data is None else data)}, data={"scope": scope})
    assert response.status_code == 201, response.text
    return response.json()


def file_id(tpl):
    with SessionLocal() as db:
        return db.get(Template, tpl["id"]).file_id


def ready(api, **kwargs):
    row = upload(api, **kwargs)
    process_file(file_id(row), WorkerAPI(api))
    return api.get(f"/v1/templates/{row['id']}").json()


def other_user():
    with SessionLocal() as db:
        user = accounts.create_user(db, "tpl-" + uuid.uuid4().hex[:12], "password-for-testing")
        token, _ = accounts.issue_session(db, user)
        db.commit()
        return {"Authorization": "Bearer " + token}


def test_registry_visibility_permissions_quota_and_library_exclusion(api):
    own = ready(api)
    company = ready(api, scope="company", name="company.pptx")
    fid = file_id(own)
    outsider = other_user()
    assert api.get(f"/v1/templates/{own['id']}", headers=outsider).status_code == 404
    assert api.delete(f"/v1/templates/{own['id']}", headers=outsider).status_code == 404
    for method in ("patch", "delete"):
        kwargs = {"json": {"name": "bad"}} if method == "patch" else {}
        assert getattr(api, method)(f"/v1/templates/{company['id']}", headers=outsider, **kwargs).status_code == 403
    assert api.post("/v1/templates", headers=outsider, files={"upload": ("a.pptx", pptx())}, data={"scope": "company"}).status_code == 403
    result = api.get("/v1/templates", headers=outsider).json()
    assert result["personal"] == [] and any(t["id"] == company["id"] and not t["can_edit"] for t in result["company"])
    assert own["status"] == "ready" and own["slides"] == 1 and own["layouts"] == 11
    assert own["thumbnail_url"] == f"/v1/templates/{own['id']}/thumbnail"
    thumb = api.get(own["thumbnail_url"])
    assert thumb.status_code == 200 and thumb.headers["content-type"] == "image/jpeg"
    assert "immutable" in thumb.headers["cache-control"]
    assert api.get(own["thumbnail_url"], headers={"If-None-Match": thumb.headers["etag"]}).status_code == 304
    assert api.get(own["thumbnail_url"], headers=outsider).status_code == 404
    assert api.get(f"/v1/files/{fid}/content").status_code == 404
    assert api.post("/v1/files/batch", json={"action": "trash", "ids": [fid]}).status_code == 404
    assert fid not in [f["id"] for f in api.get("/v1/files").json()]
    _, search_run = start(api)
    search = api.post(f"/internal/runs/{search_run['id']}/library/search", headers=SERVICE, json={"query": "demo"})
    assert search.status_code == 200 and search.json()["files"] == []
    assert api.post(f"/internal/runs/{search_run['id']}/library/read", headers=SERVICE, json={"file_id": fid}).status_code == 404
    assert api.post("/v1/shares", json={"kind": "file", "file_id": fid, "audience": "everyone"}).status_code == 404
    with SessionLocal() as db:
        file = db.get(FileObject, fid)
        assert library_bytes(db, file.user_id) == file.size
        assert (file.meta or {}).get("summary") is None
    assert api.patch(f"/v1/templates/{own['id']}", json={"name": " renamed "}).json()["name"] == "renamed"
    assert api.delete(f"/v1/templates/{own['id']}").status_code == 204
    assert api.get(f"/v1/templates/{own['id']}").status_code == 404
    with SessionLocal() as db:
        assert db.get(FileObject, fid).purged_at is None


@pytest.mark.parametrize(
    "name,media",
    [
        ("a.pptm", "application/octet-stream"),
        ("a.txt", "text/plain"),
        ("a.pptx", "application/vnd.ms-powerpoint.presentation.macroEnabled.12"),
    ],
)
def test_reject_types(api, name, media):
    response = api.post("/v1/templates", files={"upload": (name, b"x", media)})
    assert response.status_code == 400


def test_limits_names_size_and_bounded_lists(api, monkeypatch):
    monkeypatch.setattr(settings, "max_personal_templates", 2)
    monkeypatch.setattr(settings, "max_company_templates", 1)
    first = upload(api)
    second = upload(api, name="second.pptx")
    assert [r["id"] for r in api.get("/v1/templates").json()["personal"]] == [second["id"], first["id"]]
    assert api.post("/v1/templates", files={"upload": ("third.pptx", pptx())}).status_code == 400
    upload(api, scope="company")
    assert api.post("/v1/templates", files={"upload": ("fourth.pptx", pptx())}, data={"scope": "company"}).status_code == 400
    api.delete(f"/v1/templates/{first['id']}")
    upload(api, name="third.pptx")
    monkeypatch.setattr(settings, "max_template_bytes", 3)
    assert api.post("/v1/templates", files={"upload": ("big.pptx", b"xxxx")}).status_code == 413
    assert api.post("/v1/templates", files={"upload": ("empty.pptx", b"")}).status_code == 400
    assert api.patch(f"/v1/templates/{second['id']}", json={"name": " "}).status_code == 422
    assert api.patch(f"/v1/templates/{second['id']}", json={"name": "x" * 61}).status_code == 422
    assert api.get("/v1/templates?kind=docx").status_code == 422


@pytest.mark.parametrize("slides,potx", [(0, False), (0, True), (1, True)])
def test_analysis_potx_and_empty_deck_thumbnail(api, slides, potx):
    row = ready(api, name="layout.potx" if potx else "empty.pptx", data=pptx(slides, potx=potx))
    assert row["status"] == "ready" and row["slides"] == slides
    assert api.get(row["thumbnail_url"]).status_code == 200
    with SessionLocal() as db:
        meta = db.get(FileObject, file_id(row)).meta
        assert meta["template"]["layout_details"][0]["placeholders"]
        assert meta["template"]["themes"][0]["fonts"]
    assert process_file(file_id(row), WorkerAPI(api))["skipped"] == "ready"


@pytest.mark.parametrize("data", [b"bad zip", pptx(macro=True)])
def test_analysis_failed_distinct_from_scanning(api, data):
    row = ready(api, data=data)
    assert row["status"] == "failed" and "解析" in row["reason"]
    assert row["thumbnail_url"] is None


def test_template_run_validation_freeze_retry_edit_and_queue(api):
    row = upload(api)
    selector = {"id": row["id"]}
    assert api.post("/v1/runs", json={"prompt": "PPT", "template": selector}).status_code == 400
    process_file(file_id(row), WorkerAPI(api))
    assert api.post("/v1/runs", json={"prompt": "PPT", "mode": "chat", "template": selector}).status_code == 400
    assert api.post("/v1/runs", json={"prompt": "PPT", "template": {**selector, "file_id": "fil_x"}}).status_code == 422
    assert api.post("/v1/runs", headers=other_user(), json={"prompt": "PPT", "template": selector}).status_code == 404
    conv = api.post("/v1/conversations", json={}).json()
    run = api.post(f"/v1/conversations/{conv['id']}/messages", json={"prompt": "PPT", "template": selector}).json()
    assert run["template"] == {"source": "template", "name": "demo"}
    frozen = run["options"]["template"]
    assert frozen["file_id"] == file_id(row)
    queued = api.post(f"/v1/conversations/{conv['id']}/queue", json={"prompt": "next PPT", "template": selector})
    assert queued.status_code == 201, queued.text
    assert queued.json()["template"] == frozen
    assert api.patch(f"/v1/conversations/{conv['id']}/queue/{queued.json()['id']}", json={"steer": True}).status_code == 400
    assert (
        api.post(f"/v1/conversations/{conv['id']}/queue", json={"prompt": "steer", "steer": True, "template": selector}).status_code == 400
    )
    api.patch(f"/v1/templates/{row['id']}", json={"name": "changed"})
    api.delete(f"/v1/templates/{row['id']}")
    claim = attempt_headers(api, run["id"])
    assert api.post(f"/internal/runs/{run['id']}/status", headers=claim, json={"status": "running"}).status_code == 200
    assert api.post(f"/internal/runs/{run['id']}/status", headers=claim, json={"status": "succeeded"}).status_code == 200
    with SessionLocal() as db:
        runs = db.scalars(select(Run).where(Run.conversation_id == conv["id"]).order_by(Run.created_at)).all()
        assert len(runs) == 2 and runs[-1].options["template"] == frozen
        next_id = runs[-1].id
    claim = attempt_headers(api, next_id)
    api.post(f"/internal/runs/{next_id}/status", headers=claim, json={"status": "running"})
    api.post(f"/internal/runs/{next_id}/status", headers=claim, json={"status": "failed"})
    retry = api.post(f"/v1/runs/{next_id}/retry", json={})
    assert retry.status_code == 201, retry.text
    assert retry.json()["options"]["template"] == frozen
    rid = retry.json()["id"]
    api.post(f"/v1/runs/{rid}/cancel")
    edit = api.post(f"/v1/runs/{rid}/edit", json={"prompt": "edited PPT"})
    assert edit.status_code == 201, edit.text
    assert edit.json()["options"]["template"] == frozen
    assert api.get(f"/internal/runs/{edit.json()['id']}", headers=SERVICE).json()["options"]["template"] == frozen


def test_one_off_file_validation_and_library_slides_filter(api, monkeypatch):
    project = api.post("/v1/projects", json={"name": "slides"}).json()
    files = []
    for name in ("UPPER.PPTX", "theme.PoTx", "not.txt"):
        response = api.post("/v1/files", files={"upload": (name, pptx())}, params={"project_id": project["id"]})
        assert response.status_code == 201, response.text
        fid = response.json()["id"]
        files.append(fid)
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            file.status, file.meta = "ready", {"scan": {"ok": True}}
            db.commit()
    first = api.get("/v1/files", params={"category": "slides", "limit": 1}).json()
    second = api.get("/v1/files", params={"category": "slides", "limit": 1, "cursor": first[0]["id"]}).json()
    assert {first[0]["id"], second[0]["id"]} == set(files[:2])
    folders = api.get("/v1/files/folders?category=slides").json()
    assert any(f["id"] == project["id"] and f["files"] == 2 for f in folders)
    selected = {"file_id": files[0]}
    out = api.post("/v1/runs", json={"prompt": "PPT", "template": selected})
    assert out.status_code == 201, out.text
    assert out.json()["template"] == {"source": "file", "name": "UPPER"}
    assert api.post("/v1/runs", headers=other_user(), json={"prompt": "PPT", "template": selected}).status_code == 400
    assert api.post("/v1/runs", json={"prompt": "PPT", "template": {"file_id": files[2]}}).status_code == 400
    monkeypatch.setattr(settings, "max_template_bytes", 1)
    assert api.post("/v1/runs", json={"prompt": "PPT", "template": selected}).status_code == 400
    monkeypatch.setattr(settings, "max_template_bytes", 30 * 1024 * 1024)
    with SessionLocal() as db:
        db.get(FileObject, files[0]).meta = {"scan": {"ok": False}}
        db.commit()
    assert api.post("/v1/runs", json={"prompt": "PPT", "template": selected}).status_code == 400


def test_broker_template_injection_and_reset(monkeypatch):
    class API(_ApiStub):
        selected = {"file_id": "template", "source": "template", "id": "tpl_a", "name": "Frozen"}

        def get_run(self, run_id):
            return {**super().get_run(run_id), "options": {"mode": "work", "template": self.selected}}

    api = API({"template": {"status": "ready", "size": 4, "bytes": b"test", "meta": {"template": {"slides": 1}}}}, [])
    manager, provider, _ = _manager_with(api)
    commands = []
    original = manager._provider_call

    def call(method, *args, **kwargs):
        if method == "exec":
            commands.append(args[1])
        return original(method, *args, **kwargs)

    monkeypatch.setattr(manager, "_provider_call", call)
    lease = manager.create("template-run")
    assert provider.get_file(lease["sandbox_id"], "/workspace/.template/template.pptx") == b"test"
    assert json.loads(provider.get_file(lease["sandbox_id"], "/workspace/.template/analysis.json")) == {"slides": 1}
    assert "rm -rf /workspace/.template" in commands
    manager.release(lease["lease_id"])
    commands.clear()
    api.selected = None
    lease = manager.create("plain-run")
    assert "rm -rf /workspace/.template" in commands
    manager.release(lease["lease_id"])
    api.selected = {"file_id": "template"}
    monkeypatch.setattr(settings, "max_template_bytes", 1)
    from ailab_platform.broker.service import LeaseError

    with pytest.raises(LeaseError):
        manager.create("oversize")


def test_internal_template_grant_is_run_bound(api):
    row = ready(api)
    response = api.post("/v1/runs", json={"prompt": "PPT", "template": {"id": row["id"]}}).json()
    _, plain = start(api)
    broker = ServicePrincipal("broker", frozenset({"files:read_for_run"}))
    with SessionLocal() as db:
        assert _bound_file(db, file_id(row), broker, response["id"], "template").id == file_id(row)
        from fastapi import HTTPException

        with pytest.raises(HTTPException):
            _bound_file(db, file_id(row), broker, plain["id"], "template")
        with pytest.raises(HTTPException):
            _bound_file(db, file_id(row), broker, response["id"], "processing")


def test_generator_preserves_template_selects_layouts_removes_examples_no_brand_logo(tmp_path, monkeypatch):
    script = Path(__file__).resolve().parents[2] / "images/sbx-office/skills/pptx/scripts/build.py"
    local = tmp_path / "build.py"
    shutil.copyfile(script, local)
    spec = importlib.util.spec_from_file_location("template_build", local)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    template_dir, brand_dir = tmp_path / "template", tmp_path / "brand"
    template_dir.mkdir()
    brand_dir.mkdir()
    (template_dir / "template.pptx").write_bytes(pptx(potx=True))
    (brand_dir / "brand.json").write_text(json.dumps({"font": "BRAND_FONT", "colors": {"primary": "#FF0000"}}))
    from tests.test_projects_brand import png

    (brand_dir / "logo.png").write_bytes(png())
    monkeypatch.setattr(module, "TEMPLATE_DIR", template_dir)
    monkeypatch.setattr(module, "BRAND_DIR", brand_dir)
    slides = [
        {"layout": "cover", "title": "New title", "subtitle": "New subtitle"},
        {"layout": "bullets", "title": "Body", "bullets": ["One", "Two"]},
        {"layout": "two_column", "title": "Compare", "left": {"bullets": ["A"]}, "right": {"bullets": ["B"]}},
        {"layout": "cards", "title": "Fallback", "cards": [{"title": "Card", "text": "Free layout"}]},
    ]
    deck = module.Deck({"slides": slides})
    output = tmp_path / "out.pptx"
    deck.build(str(output))
    prs = Presentation(output)
    assert len(prs.slides) == 4 and len(prs.slide_masters) == 1
    assert [s.slide_layout.name for s in list(prs.slides)[:3]] == ["Title Slide", "Title and Content", "Two Content"]
    text = " ".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    assert "SAMPLE TEXT" not in text and all(word in text for word in ("New title", "One", "Two", "Compare", "Free layout"))
    assert not any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes)
    assert module.FONT != "BRAND_FONT"
    assert deck.t["primary"] != "FF0000"
    assert prs.slide_width == Inches(10)


def test_infected_template_rejected_before_analysis(api):
    row = ready(api, data=b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE")
    assert row["status"] == "rejected" and "安全扫描" in row["reason"]
    with SessionLocal() as db:
        assert "template" not in db.get(FileObject, file_id(row)).meta


def test_schedules_explicitly_reject_template(api):
    from tests.test_schedules_citations import body

    response = api.post("/v1/schedules", json={**body(), "template": {"id": "tpl_any"}})
    assert response.status_code == 422


def test_resume_and_fork_retain_frozen_selection(api):
    from ailab_platform.adapters.runtime import fresh_state
    from tests.test_runtime_protocol import pause, save

    row = ready(api)
    conv = api.post("/v1/conversations", json={}).json()
    response = api.post(f"/v1/conversations/{conv['id']}/messages", json={"prompt": "PPT", "template": {"id": row["id"]}})
    assert response.status_code == 201, response.text
    run = response.json()
    rid, frozen = run["id"], run["options"]["template"]
    headers = attempt_headers(api, rid)
    assert api.post(f"/internal/runs/{rid}/transition", headers=headers, json={"status": "running"}).status_code == 200
    state = fresh_state()
    assert save(api, rid, headers, state=state).status_code == 200
    assert pause(api, rid, headers).status_code == 200
    api.delete(f"/v1/templates/{row['id']}")
    answer = api.post(f"/v1/runs/{rid}/answer", json={"question_id": rid + ":q", "answer": "A"})
    assert answer.status_code == 202, answer.text
    assert answer.json()["options"]["template"] == frozen
    with SessionLocal() as db:
        job_id = db.get(Run, rid).runtime_meta["job_id"]
    claim = api.post(f"/internal/jobs/{job_id}/claim", headers=SERVICE, json={"consumer": "template-resume"}).json()
    assert claim["state"] == "claimed", claim
    headers = {**SERVICE, "X-Run-Attempt": claim["attempt_id"]}
    api.post(f"/internal/runs/{rid}/transition", headers=headers, json={"status": "running"})
    api.post(f"/internal/runs/{rid}/status", headers=headers, json={"status": "succeeded"})
    fork = api.post(f"/v1/conversations/{conv['id']}/fork", json={"run_id": rid, "mode": "work"})
    assert fork.status_code == 201, fork.text
    copied = api.get(f"/v1/conversations/{fork.json()['id']}").json()["turns"][-1]
    assert copied["options"]["template"] == frozen
