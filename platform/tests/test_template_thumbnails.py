"""First-page fidelity, optional rendering, and bounded renderer process trees."""

import io
import shutil
import signal
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor

from ailab_platform.config import settings
from ailab_platform.db import SessionLocal
from ailab_platform.models import FileObject
from ailab_platform.storage import get_storage
from ailab_platform.worker import isolate, pipeline, template_preview
from tests.test_projects_brand import api as api
from tests.test_templates import file_id, ready


def test_thumbnail_uses_first_rendered_page_and_downscales(tmp_path, monkeypatch):
    calls = []

    def pages(source, output, **kwargs):
        calls.append((source, kwargs))
        Image.new("RGB", (1440, 1080), "#123456").save(output / "p-01.jpg")
        Image.new("RGB", (1440, 1080), "red").save(output / "p-02.jpg")
        return {"pages": ["p-01.jpg", "p-02.jpg"]}

    monkeypatch.setattr("ailab_platform.sandbox_tools.render_preview.pages", pages)
    source = tmp_path / "source.pptx"
    assert template_preview.render(source, tmp_path) == {"thumbnail": "rendered"}
    assert calls == [(source, {"max_pages": 1, "scale_to": 480, "width_only": True})]
    with Image.open(tmp_path / "thumb.jpg") as image:
        assert image.format == "JPEG" and image.size == (480, 360)
        assert max(abs(a - b) for a, b in zip(image.getpixel((10, 10)), (18, 52, 86))) < 4


def test_shared_renderer_limits_pdf_export_and_raster_to_one_page(tmp_path, monkeypatch):
    from ailab_platform.sandbox_tools import render_preview

    conversions, commands = [], []

    def office(source, fmt, output):
        conversions.append(fmt)
        target = output / "sample.pdf"
        target.write_bytes(b"pdf")
        return target

    def run(command, **kwargs):
        commands.append(command)
        return SimpleNamespace(stdout="Pages: 1\n")

    monkeypatch.setattr(render_preview, "soffice", office)
    monkeypatch.setattr(render_preview.subprocess, "run", run)
    render_preview.pages(tmp_path / "sample.pptx", tmp_path, max_pages=1, scale_to=480, width_only=True)
    assert '"PageRange":{"type":"string","value":"1"}' in conversions[0]
    raster = commands[-1]
    assert raster[raster.index("-l") + 1] == "1"
    assert raster[raster.index("-scale-to-x") + 1] == "480"
    assert raster[raster.index("-scale-to-y") + 1] == "-1"


def test_ready_template_caches_rendered_jpeg_once(api, monkeypatch):
    original, calls = pipeline.execute, []

    def execute(operation, source, output, **kwargs):
        if operation == "template_thumbnail":
            calls.append(kwargs)
            output.mkdir(parents=True, exist_ok=True)
            Image.new("RGB", (480, 360), "#123456").save(output / "thumb.jpg")
            return {"thumbnail": "rendered"}
        return original(operation, source, output, **kwargs)

    monkeypatch.setattr(pipeline, "execute", execute)
    row = ready(api)
    response = api.get(row["thumbnail_url"])
    assert response.status_code == 200 and response.headers["content-type"] == "image/jpeg"
    assert response.content[:2] == b"\xff\xd8"
    assert "template-v2" in response.headers["etag"] and "immutable" in response.headers["cache-control"]
    assert api.get(row["thumbnail_url"], headers={"If-None-Match": response.headers["etag"]}).status_code == 304
    with SessionLocal() as db:
        meta = db.get(FileObject, file_id(row)).meta
        assert meta["template_thumbnail"] == "rendered" and meta["thumb_key"].endswith("template-v2.jpg")
    from tests.test_projects_brand import WorkerAPI

    assert pipeline.process_file(file_id(row), WorkerAPI(api))["skipped"] == "ready"
    cached = pipeline._template_thumbnail(file_id(row), None, get_storage(), None)
    assert cached["template_thumbnail"] == "cached"
    assert len(calls) == 1
    assert calls[0]["timeout"] == settings.parse_timeout_seconds * 0.85
    assert calls[0]["watch_process_tree"] is True


@pytest.mark.parametrize("failure", ["timeout", "memory_limit", "corrupt", "missing_binary"])
def test_renderer_failure_falls_back_without_rejecting_template(api, monkeypatch, failure):
    original, calls = pipeline.execute, []

    def execute(operation, source, output, **kwargs):
        calls.append(operation)
        if operation == "template_thumbnail":
            raise FileNotFoundError() if failure == "missing_binary" else isolate.IsolationError(failure)
        return original(operation, source, output, **kwargs)

    monkeypatch.setattr(pipeline, "execute", execute)
    row = ready(api)
    assert row["status"] == "ready" and row["reason"] is None
    assert calls == ["template", "template_thumbnail", "template_sketch"]
    assert api.get(row["thumbnail_url"]).headers["content-type"] == "image/jpeg"
    with SessionLocal() as db:
        assert db.get(FileObject, file_id(row)).meta["template_thumbnail"] == "sketch"


def test_both_thumbnail_paths_failing_does_not_destroy_analysis(api, monkeypatch):
    original = pipeline.execute

    def execute(operation, *args, **kwargs):
        if operation in {"template_thumbnail", "template_sketch"}:
            raise isolate.IsolationError("memory_limit")
        return original(operation, *args, **kwargs)

    monkeypatch.setattr(pipeline, "execute", execute)
    row = ready(api)
    assert row["status"] == "ready" and row["reason"] is None and row["thumbnail_url"] is None
    assert row["layouts"] == 11 and row["slides"] == 1


def test_old_webp_thumbnail_remains_readable(api, monkeypatch):
    monkeypatch.setattr(pipeline, "_template_thumbnail", lambda *args: {})
    row = ready(api)
    fid = file_id(row)
    image = io.BytesIO()
    Image.new("RGB", (30, 20), "blue").save(image, "WEBP")
    key = f"derived/{fid}/template-v1.webp"
    get_storage().put_bytes(key, image.getvalue(), "image/webp")
    with SessionLocal() as db:
        file = db.get(FileObject, fid)
        file.meta = {**file.meta, "thumb_key": key}
        db.commit()
    response = api.get(f"/v1/templates/{row['id']}/thumbnail")
    assert response.headers["content-type"] == "image/webp"
    assert "template-v1" in response.headers["etag"]


def test_renderer_memory_watch_includes_descendant_rss(monkeypatch):
    monkeypatch.setattr(isolate.subprocess, "run", lambda *a, **k: SimpleNamespace(stdout="100 1024\n100 2048\n999 4096\n"))
    assert isolate._process_group_rss(100) == 3072 * 1024


@pytest.mark.parametrize("failure", ["memory", "timeout"])
def test_limited_renderer_is_killed_on_budget_exhaustion(tmp_path, monkeypatch, failure):
    killed = []

    class Process:
        pid = 88888

        def poll(self):
            return None

        def wait(self):
            pass

    monkeypatch.setattr(isolate.subprocess, "Popen", lambda *a, **k: Process())
    monkeypatch.setattr(isolate, "_process_group_rss", lambda pid: 300 * 1024 * 1024)
    monkeypatch.setattr(isolate.os, "killpg", lambda pid, sig: killed.append((pid, sig)))
    with pytest.raises(isolate.IsolationError, match="memory_limit" if failure == "memory" else "timeout"):
        isolate.execute(
            "template_thumbnail",
            tmp_path / "x.pptx",
            tmp_path / "out",
            watch_process_tree=True,
            timeout=60 if failure == "memory" else 1e-12,
        )
    assert killed == [(88888, signal.SIGKILL)]


@pytest.mark.skipif(not all(shutil.which(tool) for tool in ("soffice", "pdfinfo", "pdftoppm")), reason="Office/Poppler not installed")
def test_real_office_thumbnail_preserves_slide_background_in_existing_256m_profile(tmp_path, monkeypatch):
    # macOS Office needs more than the default 192 MiB; do not silently raise the
    # production default. Verify the existing, configurable upper profile here.
    monkeypatch.setattr(settings, "parse_memory_mb", 256)
    source = tmp_path / "faithful.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Actual slide"
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("123456")
    prs.save(source)
    output = tmp_path / "out"
    result = isolate.execute(
        "template_thumbnail", source, output, watch_process_tree=True, measure=True, timeout=settings.parse_timeout_seconds * 0.85
    )
    assert result["thumbnail"] == "rendered"
    assert result["measurement"]["peak_process_tree_rss_bytes"] <= settings.parse_memory_mb * 1024 * 1024
    with Image.open(output / "thumb.jpg") as image:
        assert image.format == "JPEG" and image.width == 480
        assert max(abs(a - b) for a, b in zip(image.getpixel((5, 5)), (18, 52, 86))) < 5


def test_renderer_exit_still_cleans_up_descendants(tmp_path, monkeypatch):
    killed = []

    class Process:
        pid = 88889

        def poll(self):
            return 0

        def wait(self):
            pass

    output = tmp_path / "out"
    output.mkdir()
    (output / "result.json").write_text('{"thumbnail":"rendered"}')
    monkeypatch.setattr(isolate.subprocess, "Popen", lambda *a, **k: Process())
    monkeypatch.setattr(isolate.os, "killpg", lambda pid, sig: killed.append((pid, sig)))
    assert isolate.execute("template_thumbnail", tmp_path / "x.pptx", output, watch_process_tree=True)["thumbnail"] == "rendered"
    assert killed == [(88889, signal.SIGKILL)]


def test_renderer_nested_output_is_counted(tmp_path):
    nested = tmp_path / "lo-profile"
    nested.mkdir()
    (nested / "large-file").write_bytes(b"x" * 1024)
    assert isolate._output_size(tmp_path, recursive=True) == 1024
    assert isolate._output_size(tmp_path, recursive=False) == 0
