"""Schedule ownership/transaction/DST contracts and source citation boundaries."""

import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select, text, update

from ailab_platform.api.main import app
from ailab_platform.api.routes_schedules import dispatch_schedules, next_occurrence
from ailab_platform.db import SessionLocal
from ailab_platform.models import AuditLog, Conversation, FileObject, OutboxMessage, Run, ScheduledTask, User, utcnow
from ailab_platform.storage import get_storage
from tests.run_claims import SERVICE
from tests.test_conversations import _finish, _other_user
from tests.test_migrations import engine as engine
from tests.test_ocr import api as api
from tests.test_ocr import rapid_upstream as rapid_upstream
from tests.test_ocr import upstream as upstream


def as_utc(value):
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def body(**kw):
    return {"title": "每日摘要", "prompt": "整理项目", "schedule": {"kind": "daily", "time": "09:00"}, "timezone": "Asia/Hong_Kong", **kw}


def test_schedule_crud_limit_ownership_and_validation():
    with TestClient(app) as c:
        auth = _other_user(c)
        other = _other_user(c)
        for bad in [
            {"timezone": "not/a/zone"},
            {"title": "x" * 61},
            {"mode": "invalid"},
            {"schedule": {"kind": "weekly", "time": "09:00", "days": []}},
            {"model": "missing"},
            {"skill": "missing"},
        ]:
            assert c.post("/v1/schedules", headers=auth, json=body(**bad)).status_code == 422
        created = c.post("/v1/schedules", headers=auth, json=body(enabled=False))
        assert created.status_code == 201, created.text
        task = created.json()
        assert task["last_status"] is None and task["next_run_at"].endswith("Z")
        for method in ["patch", "delete", "post"]:
            url = "/v1/schedules/" + task["id"] + ("/run" if method == "post" else "")
            kwargs = {"json": {"title": "foreign"}} if method == "patch" else {}
            assert getattr(c, method)(url, headers=other, **kwargs).status_code == 404
        assert c.get("/v1/schedules", headers=other).json() == []
        patched = c.patch("/v1/schedules/" + task["id"], headers=auth, json={"title": "new", "enabled": True})
        assert patched.status_code == 200 and patched.json()["title"] == "new"
        assert c.patch("/v1/schedules/" + task["id"], headers=auth, json={"timezone": None}).status_code == 422
        with ThreadPoolExecutor(max_workers=8) as pool:
            statuses = list(pool.map(lambda _: c.post("/v1/schedules", headers=auth, json=body()).status_code, range(8)))
        assert statuses.count(201) == 4 and statuses.count(422) == 4
        assert len(c.get("/v1/schedules", headers=auth).json()) == 5
        assert c.post("/v1/schedules", headers=auth, json=body()).json()["detail"] == "最多 5 个定时任务"
        assert c.delete("/v1/schedules/" + task["id"], headers=auth).status_code == 204
        with SessionLocal() as db:
            actions = set(db.scalars(select(AuditLog.action).where(AuditLog.target_id == task["id"])))
            assert {"schedule.created", "schedule.updated", "schedule.deleted"} <= actions


def test_schedule_run_now_preserves_next_and_tracks_terminal():
    with TestClient(app) as c:
        auth = _other_user(c)
        project = c.post("/v1/projects", headers=auth, json={"name": uuid.uuid4().hex}).json()
        task = c.post("/v1/schedules", headers=auth, json=body(project_id=project["id"])).json()
        result = c.post(f"/v1/schedules/{task['id']}/run", headers=auth)
        assert result.status_code == 201, result.text
        run = result.json()
        with SessionLocal() as db:
            conv = db.get(Conversation, run["conversation_id"])
            assert conv.project_id == project["id"] and conv.title.startswith("每日摘要 · ")
            assert db.scalar(select(func.count()).select_from(OutboxMessage).where(OutboxMessage.id == run["run_id"])) == 1
        _finish(c, run["run_id"], "finished")
        listed = c.get("/v1/schedules", headers=auth).json()[0]
        assert listed["last_status"] == "succeeded" and listed["next_run_at"] == task["next_run_at"]
        assert c.delete("/v1/projects/" + project["id"], headers=auth).status_code == 204
        listed = c.get("/v1/schedules", headers=auth).json()[0]
        assert listed["project_id"] is None and not listed["enabled"]


def test_due_claim_coalesces_with_two_dispatchers_and_records_failures():
    with TestClient(app) as c:
        auth = _other_user(c)
        task = c.post("/v1/schedules", headers=auth, json=body()).json()
        now = utcnow()
        with SessionLocal() as db:
            row = db.get(ScheduledTask, task["id"])
            row.next_run_at = now - timedelta(days=10)
            owner = row.user_id
            db.commit()
        with ThreadPoolExecutor(max_workers=2) as pool:
            counts = list(pool.map(lambda _: dispatch_schedules(now), range(2)))
        assert sum(counts) == 1
        with SessionLocal() as db:
            row = db.get(ScheduledTask, task["id"])
            assert as_utc(row.next_run_at) > now and row.last_run_id
            assert db.scalar(select(func.count()).select_from(Conversation).where(Conversation.user_id == owner)) == 1
            db.get(Run, row.last_run_id).status = "succeeded"
            row.next_run_at = now - timedelta(seconds=1)
            row.skill = "missing"
            db.commit()
        assert dispatch_schedules(now) == 1
        with SessionLocal() as db:
            row = db.get(ScheduledTask, task["id"])
            assert row.last_status == "failed" and row.last_error and row.last_run_id is None
            assert db.scalar(select(func.count()).select_from(Conversation).where(Conversation.user_id == owner)) == 1
            row.next_run_at = now - timedelta(seconds=1)
            row.skill = None
            db.get(User, owner).disabled_at = now
            db.commit()
        assert dispatch_schedules(now) == 1
        with SessionLocal() as db:
            assert db.get(ScheduledTask, task["id"]).last_error == "账号已停用"
            assert db.scalar(select(func.count()).select_from(Run).where(Run.user_id == owner)) == 1
        assert dispatch_schedules(now) == 0


def test_dst_nonexistent_and_repeated_wall_times_and_weekdays():
    daily = {"kind": "daily", "time": "02:30"}
    assert next_occurrence(daily, "America/New_York", datetime(2026, 3, 8, 5, tzinfo=UTC)) == datetime(2026, 3, 9, 6, 30, tzinfo=UTC)
    daily["time"] = "01:30"
    first = next_occurrence(daily, "America/New_York", datetime(2026, 11, 1, 4, tzinfo=UTC))
    assert first == datetime(2026, 11, 1, 5, 30, tzinfo=UTC)
    assert next_occurrence(daily, "America/New_York", first) == datetime(2026, 11, 2, 6, 30, tzinfo=UTC)
    assert next_occurrence({"kind": "weekdays", "time": "09:00"}, "UTC", datetime(2026, 10, 23, 10, tzinfo=UTC)).weekday() == 0
    assert next_occurrence({"kind": "weekly", "time": "09:00", "days": [2]}, "UTC", datetime(2026, 10, 23, tzinfo=UTC)).weekday() == 2


def test_schedule_due_query_index():
    with TestClient(app), SessionLocal() as db:
        sql = "SELECT id FROM scheduled_tasks WHERE enabled = true AND next_run_at <= '2026-01-01' ORDER BY next_run_at LIMIT 100"
        if db.bind.dialect.name == "postgresql":
            db.execute(text("SET LOCAL enable_seqscan = off"))
            plan = db.execute(text("EXPLAIN " + sql))
        else:
            plan = db.execute(text("EXPLAIN QUERY PLAN " + sql))
        assert "ix_scheduled_tasks_due" in str(list(plan))


def test_library_page_windows_and_project_bootstrap():
    from ailab_platform.agent.prompts import project_library_prompt, system_prompt_for

    with TestClient(app) as c:
        auth = _other_user(c)
        project = c.post("/v1/projects", headers=auth, json={"name": uuid.uuid4().hex}).json()
        task = c.post("/v1/schedules", headers=auth, json=body(project_id=project["id"])).json()
        run = c.post(f"/v1/schedules/{task['id']}/run", headers=auth).json()
        with SessionLocal() as db:
            file = FileObject(
                user_id=task["user_id"],
                project_id=project["id"],
                filename="报告.pdf",
                storage_key="unused",
                status="ready",
                meta={"text_key": "derived/" + uuid.uuid4().hex},
            )
            db.add(file)
            db.commit()
            file_id, key = file.id, file.meta["text_key"]
        content = "\f[page:1]\n" + "中" * 32000 + "\f[page:2]\nsecond"
        get_storage().put_bytes(key, content.encode())
        url = f"/internal/runs/{run['run_id']}"
        first = c.post(url + "/library/read", headers=SERVICE, json={"file_id": file_id}).json()
        assert first["pages"] == [{"page": 1, "offset": 10}]
        last = c.post(url + "/library/read", headers=SERVICE, json={"file_id": file_id, "offset": 30000}).json()
        assert last["pages"] == [{"page": 1, "offset": 0}, {"page": 2, "offset": 2020}]
        assert first["text"] + last["text"] == content
        bootstrap = c.get(url, headers=SERVICE).json()
        assert bootstrap["project_library"] == [{"file_id": file_id, "filename": "报告.pdf"}]
        plain = c.post("/v1/runs", headers=auth, json={"prompt": "hello"}).json()
        assert c.get(f"/internal/runs/{plain['id']}", headers=SERVICE).json()["project_library"] == []
        get_storage().put_bytes(key, b"legacy text without pages")
        assert c.post(url + "/library/read", headers=SERVICE, json={"file_id": file_id}).json()["pages"] == []
        assert file_id in project_library_prompt(bootstrap["project_library"])
        assert project_library_prompt([]) == ""
        for mode in ["chat", "work"]:
            prompt = system_prompt_for(mode)
            assert "ailab-file://<file_id>?page=N" in prompt and "不得引用未读" in prompt


def test_worker_pdf_pptx_page_markers_and_existing_limits(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter

    from ailab_platform.worker.parsers import MAX_TEXT, _docx, _pdf, _pptx, _xlsx

    pdf = tmp_path / "test.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.add_blank_page(100, 100)
    writer.write(pdf)
    assert _pdf(pdf)["text"] == "\f[page:1]\n\f[page:2]\n"
    prs = Presentation()
    for value in ["one", "two" * MAX_TEXT]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(2)).text = value
    ppt = tmp_path / "test.pptx"
    prs.save(ppt)
    result = _pptx(ppt)
    assert "\f[page:2]" in result["text"] and len(result["text"]) == MAX_TEXT and result["truncated"]
    import docx
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "收入"
    path = tmp_path / "test.xlsx"
    wb.save(path)
    assert "# 收入" in _xlsx(path)["text"]
    doc = docx.Document()
    doc.add_paragraph("text")
    path = tmp_path / "test.docx"
    doc.save(path)
    assert "[page:" not in _docx(path)["text"]  # No stable rendered pagination in the DOCX package.


def test_schedule_migration_roundtrip(engine):
    from alembic.script import ScriptDirectory
    from sqlalchemy import inspect

    from ailab_platform.migrate import alembic_config
    from tests.test_migrations import _downgrade_to, _upgrade_to

    revision = ScriptDirectory.from_config(alembic_config()).get_revision("0024")
    _upgrade_to(engine, revision.down_revision)
    _upgrade_to(engine, "0024")
    assert "scheduled_tasks" in inspect(engine).get_table_names()
    _downgrade_to(engine, revision.down_revision)
    assert "scheduled_tasks" not in inspect(engine).get_table_names()
    _upgrade_to(engine, "0024")


def test_schedule_checks_quota_capacity_and_current_skill_state(monkeypatch):
    from ailab_platform.agent.prompts import skills
    from ailab_platform.config import settings

    with TestClient(app) as c:
        auth = _other_user(c)
        skill = "pptx"
        assert skill in skills()
        task = c.post("/v1/schedules", headers=auth, json=body(skill=skill, mode="chat")).json()
        assert task["mode"] == "chat"
        with SessionLocal() as db:
            user = db.get(User, task["user_id"])
            user.settings = {**(user.settings or {}), "disabled_skills": [skill]}
            db.commit()
        assert c.post("/v1/schedules", headers=auth, json=body(skill=skill)).status_code == 422
        assert c.post(f"/v1/schedules/{task['id']}/run", headers=auth).status_code == 422
        assert c.get("/v1/schedules", headers=auth).json()[0]["last_status"] == "failed"
        c.patch(f"/v1/schedules/{task['id']}", headers=auth, json={"skill": None})
        monkeypatch.setattr(settings, "max_active_runs_per_user", 1)
        c.post("/v1/runs", headers=auth, json={"prompt": "occupy slot"})
        assert c.post(f"/v1/schedules/{task['id']}/run", headers=auth).status_code == 422
        with SessionLocal() as db:
            assert db.get(ScheduledTask, task["id"]).last_error == "额度或并发任务已达上限"
            assert db.scalar(select(func.count()).select_from(Conversation).where(Conversation.user_id == task["user_id"])) == 0
        monkeypatch.setattr(settings, "max_active_runs_per_user", 1000)
        monkeypatch.setattr(settings, "daily_token_limit", 0)
        assert c.post(f"/v1/schedules/{task['id']}/run", headers=auth).status_code == 422
        with SessionLocal() as db:
            assert db.get(ScheduledTask, task["id"]).last_error == "额度或并发任务已达上限"


def test_ocr_text_preserves_original_pdf_page_numbers(api, rapid_upstream):
    from ailab_platform.worker.ocr import process_ocr
    from ailab_platform.worker.pipeline import process_file

    # Image OCR represents the sole preview page, not an invented PDF sequence.
    from tests.test_ocr import AUTH, OCRClient, photo, upload

    fid = upload(api, name="page.png", data=photo(), mime="image/png")
    process_file(fid, platform=OCRClient(api))
    api.post(f"/v1/files/{fid}/extract-text", headers=AUTH, json={})
    process_ocr(fid, OCRClient(api))
    with SessionLocal() as db:
        key = db.get(FileObject, fid).meta["ocr_text_key"]
    assert get_storage().get_bytes(key).decode().startswith("\f[page:1]\n")


def test_dispatch_crash_rolls_back_claim_conversation_run_and_outbox(monkeypatch):
    import ailab_platform.api.routes_schedules as schedules

    with TestClient(app) as c:
        auth = _other_user(c)
        task = c.post("/v1/schedules", headers=auth, json=body()).json()
        now = utcnow()
        observed = now - timedelta(seconds=1)
        with SessionLocal() as db:
            db.get(ScheduledTask, task["id"]).next_run_at = observed
            db.commit()
        original = schedules.start_run

        def crash(*args, **kwargs):
            original(*args, **kwargs)
            raise RuntimeError("injected crash after outbox staging")

        with monkeypatch.context() as patch:
            patch.setattr(schedules, "start_run", crash)
            with pytest.raises(RuntimeError, match="injected crash"):
                dispatch_schedules(now)
        with SessionLocal() as db:
            assert as_utc(db.get(ScheduledTask, task["id"]).next_run_at) == observed
            assert not db.scalar(select(Conversation.id).where(Conversation.user_id == task["user_id"]))
            assert not db.scalar(select(Run.id).where(Run.user_id == task["user_id"]))
        assert dispatch_schedules(now) == 1
        with SessionLocal() as db:
            row = db.get(ScheduledTask, task["id"])
            assert db.get(OutboxMessage, row.last_run_id) is not None


@pytest.fixture
def guardrail_task():
    with TestClient(app) as c:
        auth = _other_user(c)
        with SessionLocal() as db:
            db.execute(update(ScheduledTask).values(enabled=False))
            db.commit()
        task = c.post("/v1/schedules", headers=auth, json=body()).json()
        yield c, auth, task


def due(task_id, now):
    with SessionLocal() as db:
        db.get(ScheduledTask, task_id).next_run_at = now
        db.commit()
    dispatch_schedules(now)


@pytest.mark.parametrize("status", ["queued", "running", "waiting_user", "interrupted"])
def test_dispatch_does_not_overlap_and_advances(guardrail_task, status):
    c, auth, task = guardrail_task
    first = c.post(f"/v1/schedules/{task['id']}/run", headers=auth).json()
    with SessionLocal() as db:
        db.get(Run, first["run_id"]).status = status
        db.commit()
    now = utcnow()
    due(task["id"], now)
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert row.last_run_id == first["run_id"] and row.failure_streak == 0
        assert row.last_status == "skipped" and row.last_error == "上一次还未结束"
        assert as_utc(row.next_run_at) > now
        assert row.recent_run_ids == [first["run_id"]]
    assert c.get("/v1/schedules", headers=auth).json()[0]["last_status"] == "skipped"


@pytest.mark.parametrize("status", ["waiting_user", "interrupted"])
def test_unanswered_occurrence_pauses_instead_of_skipping_forever(guardrail_task, status):
    from ailab_platform.models import Notification

    c, auth, task = guardrail_task
    first = c.post(f"/v1/schedules/{task['id']}/run", headers=auth).json()
    with SessionLocal() as db:
        run = db.get(Run, first["run_id"])
        run.status, run.started_at = status, utcnow() - timedelta(days=8)
        db.commit()
    due(task["id"], utcnow())
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert not row.enabled and row.paused_reason == "unread" and row.last_run_id == first["run_id"]
        assert db.scalar(select(Notification).where(Notification.user_id == row.user_id, Notification.kind == "schedule.paused"))


def test_end_date_local_validation_and_single_pause(guardrail_task, monkeypatch):
    import ailab_platform.api.routes_schedules as schedules
    from ailab_platform.models import Notification

    c, auth, task = guardrail_task
    now = datetime(2026, 11, 2, 1, tzinfo=UTC)
    monkeypatch.setattr(schedules, "utcnow", lambda: now)
    # Still November 1 in Los Angeles.
    changed = c.patch(f"/v1/schedules/{task['id']}", headers=auth, json={"ends_on": "2026-11-01", "timezone": "America/Los_Angeles"})
    assert changed.status_code == 200, changed.text
    bad = c.patch(f"/v1/schedules/{task['id']}", headers=auth, json={"ends_on": "2026-10-31"})
    assert bad.status_code == 422 and bad.json()["detail"] == "截止日期不能早于今天"
    bad = c.post("/v1/schedules", headers=auth, json=body(ends_on="2026-11-01", timezone="Asia/Shanghai"))
    assert bad.status_code == 422 and bad.json()["detail"] == "截止日期不能早于今天"
    due(task["id"], now + timedelta(days=1))
    dispatch_schedules(now + timedelta(days=2))
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert not row.enabled and row.paused_reason == "ended" and row.last_run_id is None
        notes = db.scalars(select(Notification).where(Notification.user_id == row.user_id)).all()
        assert len(notes) == 1 and notes[0].kind == "schedule.paused"
        assert notes[0].title == row.title and notes[0].conversation_id is None
    assert c.get("/v1/schedules", headers=auth).status_code == 200  # Expired dates remain serializable.
    assert c.get("/v1/notifications", headers=auth).json()["items"][0]["kind"] == "schedule.paused"


def test_final_allowed_occurrence_runs_then_ends(guardrail_task):
    c, auth, task = guardrail_task
    today = utcnow().date()
    assert c.patch(f"/v1/schedules/{task['id']}", headers=auth, json={"ends_on": today.isoformat(), "timezone": "UTC"}).status_code == 200
    due(task["id"], utcnow().replace(hour=23, minute=59))
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert row.last_run_id and not row.enabled and row.paused_reason == "ended"


@pytest.mark.parametrize("seen", [None, "read", "deleted", "canceled", "no_notification", "young", "two"])
def test_unattended_pause_requires_three_unseen_and_seven_days(guardrail_task, seen):
    from ailab_platform.models import Notification

    _, _, task = guardrail_task
    now = utcnow()
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        ids = []
        for i in range(2 if seen == "two" else 3):
            conv = Conversation(user_id=row.user_id, title="occurrence", deleted_at=now if seen == "deleted" and i == 0 else None)
            db.add(conv)
            db.flush()
            run = Run(
                user_id=row.user_id,
                conversation_id=conv.id,
                prompt="p",
                model="m",
                status="canceled" if seen == "canceled" and i == 0 else "succeeded",
                started_at=now - timedelta(days=(6 if seen == "young" else 7) - i),
            )
            db.add(run)
            db.flush()
            ids.append(run.id)
            if not (seen == "no_notification" and i == 0):
                db.add(
                    Notification(
                        user_id=row.user_id,
                        run_id=run.id,
                        conversation_id=conv.id,
                        kind="schedule.succeeded",
                        title=row.title,
                        read_at=now if seen == "read" and i == 0 else None,
                    )
                )
        row.recent_run_ids, row.last_run_id = ids, ids[-1]
        db.commit()
    due(task["id"], now)
    dispatch_schedules(now)
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        pauses = db.scalars(select(Notification).where(Notification.user_id == row.user_id, Notification.kind == "schedule.paused")).all()
        if seen is None:
            assert not row.enabled and row.paused_reason == "unread" and row.last_run_id == ids[-1]
            assert len(pauses) == 1
        else:
            assert row.enabled and row.paused_reason is None and row.last_run_id not in ids
            assert len(row.recent_run_ids) == 3 and not pauses


def test_dispatch_start_failures_pause_but_run_now_failure_does_not(guardrail_task):
    from ailab_platform.models import Notification

    c, auth, task = guardrail_task
    with SessionLocal() as db:
        db.get(ScheduledTask, task["id"]).skill = "missing"
        db.commit()
    assert c.post(f"/v1/schedules/{task['id']}/run", headers=auth).status_code == 422
    for i in range(3):
        due(task["id"], utcnow() + timedelta(days=i))
        with SessionLocal() as db:
            row = db.get(ScheduledTask, task["id"])
            assert row.failure_streak == i + 1
            assert row.enabled == (i < 2)
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert row.paused_reason == "failures" and row.recent_run_ids == []
        assert db.scalar(select(func.count()).select_from(Notification).where(Notification.user_id == row.user_id)) == 1


def test_terminal_failure_streak_success_reset_and_late_completion_idempotency(guardrail_task):
    from ailab_platform.models import Notification
    from ailab_platform.run_state import transition_run

    c, auth, task = guardrail_task
    for status in ["failed", "succeeded", "failed", "failed", "failed"]:
        rid = c.post(f"/v1/schedules/{task['id']}/run", headers=auth).json()["run_id"]
        with SessionLocal() as db:
            if status == "succeeded":
                transition_run(db, rid, "running", commit=False)
            transition_run(db, rid, status, commit=False)
            transition_run(db, rid, status, commit=False)
            db.commit()
            row = db.get(ScheduledTask, task["id"])
            if status == "succeeded":
                assert row.failure_streak == 0
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert row.failure_streak == 3 and row.paused_reason == "failures" and not row.enabled
        assert len(row.recent_run_ids) == 3
        assert (
            db.scalar(
                select(func.count())
                .select_from(Notification)
                .where(Notification.user_id == row.user_id, Notification.kind == "schedule.paused")
            )
            == 1
        )


@pytest.mark.parametrize("patch", [{"enabled": True}, {"schedule": {"kind": "daily", "time": "10:00"}}, {"ends_on": None}])
def test_edit_or_reenable_resets_pause_history(guardrail_task, patch):
    c, auth, task = guardrail_task
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        row.enabled, row.paused_reason, row.failure_streak, row.recent_run_ids = False, "unread", 2, ["a", "b", "c"]
        db.commit()
    response = c.patch(f"/v1/schedules/{task['id']}", headers=auth, json=patch)
    assert response.status_code == 200 and response.json()["paused_reason"] is None
    with SessionLocal() as db:
        row = db.get(ScheduledTask, task["id"])
        assert row.recent_run_ids == [] and row.failure_streak == 0
