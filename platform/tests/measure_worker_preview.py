"""Reproduce W1's local fixture/child measurements, without a developer environment/database.

AILAB_ENV_FILE=/tmp/no-env AILAB_ENV=test .venv/bin/python -m tests.measure_worker_preview
"""

import io
import json
import os
import platform
import random
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook
from PIL import Image
from pptx import Presentation


def pad_zip(source, target_size):
    missing = target_size - source.stat().st_size
    if missing > 128:
        with zipfile.ZipFile(source, "a", zipfile.ZIP_STORED) as archive:
            archive.writestr("benchmark-padding.bin", os.urandom(missing - 128))


def main():
    with tempfile.TemporaryDirectory(prefix="ailab-w1-measure-") as tmp:
        root = Path(tmp)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        # The 20 MiB file contains an embedded incompressible image, not only a ZIP padding entry.
        image = io.BytesIO()
        Image.frombytes("RGB", (2600, 2600), random.Random(42).randbytes(2600 * 2600 * 3)).save(image, "PNG")
        image.seek(0)
        slide.shapes.add_picture(image, 0, 0)
        pptx = root / "20mb.pptx"
        deck.save(pptx)
        pad_zip(pptx, 20 * 1024 * 1024)
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("data")
        rng = random.Random(42)
        for _ in range(12_000):
            ws.append([rng.randbytes(64).hex() for _ in range(10)])
        xlsx = root / "10mb.xlsx"
        wb.save(xlsx)
        pad_zip(xlsx, 10 * 1024 * 1024)
        results = {"host": platform.platform(), "fixtures": []}
        for op, source in [("extract", pptx), ("sheets", xlsx)]:
            start = time.monotonic()
            runner = """
import json, sys, resource
from pathlib import Path
import ailab_platform.worker.main
from ailab_platform.worker.isolate import execute
result = execute(sys.argv[1], Path(sys.argv[2]), Path(sys.argv[3]), filename=Path(sys.argv[2]).name, measure=True, timeout=120)
result["measurement"]["worker_parent_peak_rss_bytes"] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
print(json.dumps(result))
"""
            process = subprocess.Popen([sys.executable, "-c", runner, op, str(source), str(root / op)], stdout=subprocess.PIPE, text=True)
            combined_peak = 0
            while process.poll() is None:
                snapshot = subprocess.run(["ps", "-axo", "pid=,ppid=,rss="], capture_output=True, text=True, check=True)
                rows = [list(map(int, line.split())) for line in snapshot.stdout.splitlines()]
                descendants = {process.pid}
                for _ in range(3):
                    descendants.update(pid for pid, parent, _rss in rows if parent in descendants)
                combined_peak = max(combined_peak, sum(rss * 1024 for pid, _, rss in rows if pid in descendants))
                time.sleep(0.03)
            stdout, _ = process.communicate(timeout=1)
            if process.returncode:
                raise RuntimeError("benchmark child failed")
            result = json.loads(stdout)
            record = {
                "operation": op,
                "bytes": source.stat().st_size,
                "seconds": round(time.monotonic() - start, 3),
                "kind": result.get("kind"),
                "sampled_worker_tree_peak_rss_bytes": combined_peak,
                **result["measurement"],
            }
            results["fixtures"].append(record)
        from docx import Document

        from ailab_platform.broker.render import LibraryRenderer
        from ailab_platform.config import settings
        from tests.test_workspace_previews import _manager

        settings.data_dir = root / "sandbox-data"
        document = Document()
        document.add_paragraph("Library preview latency measurement")
        data = io.BytesIO()
        document.save(data)
        manager, api, storage = _manager(root)
        api.get_file = lambda fid: {
            "status": "ready",
            "filename": "sample.docx",
            "meta": {"preview_job_id": "measurement", "preview": {"status": "pending"}},
        }
        api.get_file_bytes = lambda fid: data.getvalue()
        renderer = LibraryRenderer(manager)
        started = time.monotonic()
        rendered = renderer.render("sample", "measurement")
        results["docx_mock_render"] = {
            "seconds": round(time.monotonic() - started, 3),
            "pages": rendered["preview"]["page_count"],
            "output_bytes": sum(storage.size(a["key"]) for a in rendered["assets"].values()),
            "active_leases_after": len(manager.active_leases()),
        }
        print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
