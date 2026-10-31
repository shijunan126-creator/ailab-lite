"""Cross-format visual coverage, OCR reuse and durable dependency recovery."""

import asyncio
import io
import time

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, text

from ailab_platform.api.main import app
from ailab_platform.catalog import ModelSpec, get_catalog
from ailab_platform.config import settings
from ailab_platform.db import SessionLocal
from ailab_platform.models import FileObject, OcrPageResult, OutboxMessage, Run
from ailab_platform.storage import get_storage
from tests.test_document_context import upload
from tests.test_file_previews import AUTH, SERVICE
from tests.test_ocr import photo


def office_image(kind, with_text=False):
    data = io.BytesIO()
    image = io.BytesIO(photo())
    if kind == "pdf":
        from PIL import Image

        Image.open(image).convert("RGB").save(data, "PDF")
    elif kind == "docx":
        from docx import Document

        doc = Document()
        if with_text:
            doc.add_paragraph("An architecture diagram follows.")
        doc.add_picture(image)
        doc.save(data)
    elif kind == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_picture(image, Inches(1), Inches(1))
        if with_text:
            slide.shapes.add_textbox(Inches(0), Inches(0), Inches(5), Inches(1)).text = "Architecture"
        doc.save(data)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image

        doc = Workbook()
        doc.active.add_image(Image(image))
        if with_text:
            doc.active["A1"] = "Architecture"
        doc.save(data)
    return data.getvalue()


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx", "xlsx"])
@pytest.mark.parametrize("with_text", [False, True])
def test_office_images_are_detected_and_never_presented_as_complete_text(kind, with_text):
    from ailab_platform.api.document_content import manifest

    with TestClient(app) as c:
        fid = upload(c, "diagram." + kind, office_image(kind, with_text), "application/" + kind)
        with SessionLocal() as db:
            doc = manifest(db.get(FileObject, fid))
            assert doc["visual_units"]
            assert all(u["quality"] == "no_text" for u in doc["units"]) if not with_text or kind == "pdf" else doc["chars"] > 0
        run = c.post("/v1/runs", headers=AUTH, json={"prompt": "解释图形关系", "file_ids": [fid]}).json()
        read = c.post(f"/internal/runs/{run['id']}/library/read", headers=SERVICE, json={"file_id": fid}).json()
        assert read["visual"]["status"] == "present"
        assert "图" in read["visual_note"]
        assert "(image)" not in read["text"] if not with_text else True


def cache_page(db, file, page, content):
    key = f"derived/{file.id}/regression/{page}.txt"
    get_storage().put_bytes(key, content.encode(), "text/plain")
    db.add(
        OcrPageResult(
            file_id=file.id,
            page=page,
            fingerprint=settings.ocr_engine_fingerprint or "test",
            generation="cached",
            text_key=key,
            engine="rapidocr:test",
            backend="cpu",
            seconds=0,
        )
    )
    db.flush()


def test_repeated_pdf_extract_returns_cached_text_without_jobs_and_legacy_placeholder_is_not_body(monkeypatch):
    monkeypatch.setattr(settings, "ocr_executor_url", "")  # cache remains readable while executor is offline
    with TestClient(app) as c:
        fid = upload(c, "scanned.pdf", office_image("pdf"), "application/pdf")
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            # Simulate the original deployed cache format exactly.
            meta = dict(file.meta)
            meta["extract"] = {**meta["extract"], "parser": "zcode-3c5b227-ailab1"}
            file.meta = meta
            from ailab_platform.models import DocumentBlock, FileDerivation

            block = db.scalar(select(DocumentBlock).join(FileDerivation).where(FileDerivation.file_id == fid))
            get_storage().put_bytes(block.text_key, b"## Page 1\n(no extractable text; the page may be scanned)", "text/plain")
            cache_page(db, file, 0, "Invoice amount 123")
            db.commit()
        run = c.post("/v1/runs", headers=AUTH, json={"prompt": "概述", "file_ids": [fid]}).json()
        endpoint = f"/internal/runs/{run['id']}/library/extract-text"
        for _ in range(3):
            result = c.post(endpoint, headers=SERVICE, json={"file_id": fid, "pages": [1]}).json()
            assert result["ocr"] == "done" and "Invoice amount 123" in result["content"]["text"]
            assert result["content"]["coverage"][0]["method"] == "ocr"
        with SessionLocal() as db:
            assert not [j for j in db.scalars(select(OutboxMessage)) if j.payload.get("file_id") == fid and j.payload.get("kind") == "ocr"]
        chat = c.post("/v1/runs", headers=AUTH, json={"prompt": "概述", "file_ids": [fid], "mode": "chat"})
        assert chat.status_code == 201
        model_input = c.get(f"/internal/runs/{chat.json()['id']}/model-input", headers=SERVICE).json()
        assert model_input["input_files"][0]["text"] == "Invoice amount 123"


def test_pdf_inflight_selection_is_not_replaced_and_failure_is_terminal(monkeypatch):
    monkeypatch.setattr(settings, "ocr_executor_url", "http://unused")
    from ailab_platform.api.routes_ocr import request_extraction

    with TestClient(app) as c:
        fid = upload(c, "scanned.pdf", office_image("pdf"), "application/pdf")
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            a = request_extraction(db, file, [1])
            generation = file.meta["ocr_generation"]
            b = request_extraction(db, file, [1])
            assert a["ocr"] == b["ocr"] == "pending" and file.meta["ocr_generation"] == generation
            file.meta = {**file.meta, "ocr": "failed", "ocr_reason": "deadline"}
            assert request_extraction(db, file, [1])["ocr"] == "failed"
            assert file.meta["ocr_generation"] == generation
            db.rollback()


def test_dependency_reconciliation_is_idempotent_bounded_and_does_not_revive_cancelled_run(monkeypatch):
    from ailab_platform.api import file_dependencies

    with TestClient(app) as c:
        fid = upload(c, "scan.pdf", office_image("pdf"), "application/pdf")
        run_id = c.post("/v1/runs", headers=AUTH, json={"prompt": "概述", "file_ids": [fid]}).json()["id"]
        with SessionLocal() as db:
            run = db.get(Run, run_id)
            run.status = "interrupted"
            run.runtime_meta = {
                **run.runtime_meta,
                "file_dependency": {"file_id": fid, "action": "extract-text", "pages": [1], "deadline": time.time() + 100},
            }
            file = db.get(FileObject, fid)
            file.meta = {**file.meta, "ocr": "pending"}
            db.commit()
            plan = select(Run.id).where(Run.status == "interrupted", Run.id > "").order_by(Run.id).limit(100)
            sql = str(plan.compile(db.bind, compile_kwargs={"literal_binds": True}))
            assert db.execute(text(("EXPLAIN QUERY PLAN " if db.bind.dialect.name == "sqlite" else "EXPLAIN ") + sql)).all()
        monkeypatch.setattr(file_dependencies, "_after", "")
        file_dependencies.reconcile()
        with SessionLocal() as db:
            assert db.get(Run, run_id).status == "interrupted"
            cache_page(db, db.get(FileObject, fid), 0, "ready OCR")
            db.commit()
        for _ in range(3):
            monkeypatch.setattr(file_dependencies, "_after", "")
            file_dependencies.reconcile()
        with SessionLocal() as db:
            assert db.get(Run, run_id).status == "queued"
            jobs = [j.id for j in db.scalars(select(OutboxMessage)) if j.payload.get("run_id") == run_id]
            assert len(jobs) == 2  # original plus one continuation
            run = db.get(Run, run_id)
            run.status = "canceled"
            run.runtime_meta = {
                **run.runtime_meta,
                "file_dependency": {"file_id": fid, "action": "extract-text", "pages": [1], "deadline": 0},
            }
            db.commit()
        monkeypatch.setattr(file_dependencies, "_after", "")
        file_dependencies.reconcile()
        with SessionLocal() as db:
            assert db.get(Run, run_id).status == "canceled"


def test_image_reference_asked_while_another_preview_is_pending_still_gets_its_pages(monkeypatch):
    from ailab_platform.api import file_dependencies
    from ailab_platform.api.routes_library_tools import document_images

    with TestClient(app) as c:
        fid = upload(c, "deck.pdf", office_image("pdf"), "application/pdf")
        run_id = c.post("/v1/runs", headers=AUTH, json={"prompt": "看图", "file_ids": [fid]}).json()["id"]
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            # the ordinary preview a delivered file gets is still running when the model asks to see the pages
            file.meta = {**file.meta, "preview": {"status": "pending", "kind": "sheets"}, "preview_job_id": "eager"}
            assert document_images(db, file, [1])["pending"]
            run = db.get(Run, run_id)
            run.status = "interrupted"
            run.runtime_meta = {
                **run.runtime_meta,
                "file_dependency": {"file_id": fid, "action": "image-reference", "pages": [1], "deadline": time.time() + 100},
            }
            db.commit()

        def renders():
            with SessionLocal() as db:
                return [j for j in db.scalars(select(OutboxMessage)) if j.payload.get("file_id") == fid and j.id != "eager"]

        before = len(renders())
        monkeypatch.setattr(file_dependencies, "_after", "")
        file_dependencies.reconcile()
        assert len(renders()) == before  # nothing stacked on the pending preview
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            file.meta = {**file.meta, "preview": {"status": "ready", "kind": "sheets"}}
            db.commit()
        for _ in range(3):
            monkeypatch.setattr(file_dependencies, "_after", "")
            file_dependencies.reconcile()
        with SessionLocal() as db:
            meta = db.get(FileObject, fid).meta
            assert meta["preview"]["status"] == "pending" and meta["ocr_render_pages"] == [1]
            assert db.get(Run, run_id).status == "interrupted"
        assert len(renders()) == before + 1  # one render, however many ticks


def test_pending_ocr_suspends_paid_tool_and_resumes_without_model_poll_or_sandbox():
    from ailab_platform.agent.loop import AgentLoop
    from ailab_platform.llm import LLMResponse, ToolCall
    from tests.test_agent_loop import FakeBroker, FakePlatform, ScriptedLLM

    class Platform(FakePlatform):
        done = False

        async def library(self, run_id, action, payload):
            return {"ocr": "done", "content": {"text": "cached words"}} if self.done else {"ocr": "pending"}

    platform, broker = Platform(), FakeBroker()
    model = ScriptedLLM(
        [LLMResponse(tool_calls=[ToolCall("ocr", "extract_text", {"file_id": "file", "pages": [1]})]), LLMResponse(content="概述")]
    )
    asyncio.run(AgentLoop(platform, broker, model).process("run_1", attempt_id="a"))
    assert platform.paused["file_dependency"]["file_id"] == "file"
    assert len(model.seen_messages) == 1 and broker.leases == 0
    platform.done = True
    asyncio.run(AgentLoop(platform, broker, model).process("run_1", attempt_id="b"))
    assert len(model.seen_messages) == 2 and broker.leases == 0
    assert platform.statuses[-1][0] == "succeeded"


def test_chat_vision_prepares_visual_pages_and_nonvision_does_not_claim_image_semantics(monkeypatch):
    from ailab_platform.models import RenderedPage

    monkeypatch.setitem(
        get_catalog().models,
        "vision-doc",
        ModelSpec(id="vision-doc", name="vision-doc", provider="mock", capabilities=["tools", "vision"], context_tokens=200000),
    )
    with TestClient(app) as c:
        fid = upload(c, "diagram.docx", office_image("docx", True), "application/docx")
        body = {"prompt": "解释架构图", "mode": "chat", "model": "vision-doc", "file_ids": [fid]}
        pending = c.post("/v1/chat/prepare", headers=AUTH, json=body)
        assert pending.status_code == 409 and pending.json()["detail"]["code"] == "chat_attachment_processing"
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            key = f"derived/{fid}/preview/p-1.jpg"
            get_storage().put_bytes(key, photo(), "image/jpeg")
            db.add(RenderedPage(file_id=fid, page=1, version="poppler-1280-v1", key=key, digest="test"))
            file.meta = {**file.meta, "ocr_preview": {"status": "ready", "page_count": 1}}
            db.commit()
        assert c.post("/v1/chat/prepare", headers=AUTH, json=body).status_code == 200
        run = c.post("/v1/runs", headers=AUTH, json=body)
        assert run.status_code == 201
        data = c.get(f"/internal/runs/{run.json()['id']}/model-input", headers=SERVICE).json()
        file = data["input_files"][0]
        assert file["visual_pages"][0]["page"] == 1 and file["visual"]["status"] == "present"
        from ailab_platform.agent.prompts import user_prompt

        text = user_prompt("解释", [{**file, "visual_pages": []}])
        assert "不能根据OCR猜测" in text and fid in text


def test_safe_file_pause_keeps_sse_live_and_quota_gate_defers_recovery(monkeypatch):
    from fastapi import HTTPException

    from ailab_platform.api import file_dependencies, routes_runtime
    from ailab_platform.api.routes_runs import _poll_events

    with TestClient(app) as c:
        fid = upload(c, "doc.txt", b"source", "text/plain")
        rid = c.post("/v1/runs", headers=AUTH, json={"prompt": "read", "file_ids": [fid]}).json()["id"]
        claimed = c.post(f"/internal/jobs/{rid}/claim", headers=SERVICE, json={"consumer": "dependency-test"}).json()
        headers = {**SERVICE, "X-Run-Attempt": claimed["attempt_id"]}
        with SessionLocal() as db:
            db.get(Run, rid).status = "running"
            db.commit()
        saved = c.put(f"/internal/runs/{rid}/checkpoint", headers=headers, json={"expected_step": 0, "state": {"runtime_version": 1}})
        assert saved.status_code == 200
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            file.meta = {**file.meta, "ocr": "pending"}
            db.commit()
        dep = {"file_id": fid, "action": "extract-text", "pages": [1], "deadline": time.time() + 100}
        paused = c.post(f"/internal/runs/{rid}/pause", headers=headers, json={"status": "interrupted", "file_dependency": dep})
        assert paused.status_code == 200
        detail = c.get(f"/v1/runs/{rid}", headers=AUTH).json()
        assert detail["result"]["file_wait"] == {"action": "extract-text"} and not detail["resumable"]
        assert _poll_events(rid, 0)[2] is True
        with SessionLocal() as db:
            cache_page(db, db.get(FileObject, fid), 0, "ready")
            db.commit()

        def full(*args):
            raise HTTPException(429, "too many active runs")

        with monkeypatch.context() as patch:
            patch.setattr(routes_runtime, "_admit_continuation", full)
            patch.setattr(file_dependencies, "_after", "")
            file_dependencies.reconcile()
        with SessionLocal() as db:
            run = db.get(Run, rid)
            assert run.status == "interrupted" and run.runtime_meta["file_dependency"]["retry_at"] > time.time()
            run.runtime_meta = {**run.runtime_meta, "file_dependency": dep}
            db.commit()
        monkeypatch.setattr(file_dependencies, "_after", "")
        file_dependencies.reconcile()
        with SessionLocal() as db:
            assert db.get(Run, rid).status == "queued"


def test_unscanned_legacy_pending_ocr_is_terminal_without_weakening_scan_gate():
    with TestClient(app) as c:
        fid = upload(c, "source.txt", b"content", "text/plain")
        with SessionLocal() as db:
            file = db.get(FileObject, fid)
            file.meta = {"ocr": "pending"}
            db.commit()
        response = c.post(f"/internal/files/{fid}/ocr/prepare", headers=SERVICE)
        assert response.status_code == 200 and response.json()["ocr"] == "failed"
        with SessionLocal() as db:
            assert db.get(FileObject, fid).meta["ocr_reason"] == "scan_required"


def test_chat_document_visual_bytes_are_required_before_any_model_call(monkeypatch):
    from ailab_platform.agent.loop import AgentLoop
    from ailab_platform.agent.tools import ToolError
    from ailab_platform.clients import ClientError

    monkeypatch.setitem(
        get_catalog().models, "vision-bytes", ModelSpec(id="vision-bytes", name="Vision", provider="mock", capabilities=["vision"])
    )

    class Platform:
        async def get_file_page_image(self, *args):
            raise ClientError("GET", "/missing", 404, "missing")

    run = {
        "prompt": "看图",
        "input_files": [
            {
                "filename": "diagram.pdf",
                "file_id": "f",
                "visual_pages": [{"id": "f", "filename": "diagram.pdf", "page": 1, "content_type": "image/jpeg"}],
            }
        ],
    }
    with pytest.raises(ToolError, match="不完整文档"):
        asyncio.run(AgentLoop(platform=Platform())._user_content("run", run, "vision-bytes"))


def test_ocr_rejects_nonexistent_image_and_document_pages_before_staging():
    from fastapi import HTTPException

    from ailab_platform.api.routes_ocr import request_extraction

    with TestClient(app) as c:
        image = upload(c, "image.png", photo(), "image/png")
        pdf = upload(c, "scan.pdf", office_image("pdf"), "application/pdf")
        with SessionLocal() as db:
            for fid in (image, pdf):
                with pytest.raises(HTTPException) as exc:
                    request_extraction(db, db.get(FileObject, fid), [2])
                assert exc.value.status_code == 422


def test_parser_upgrade_reuses_private_ocr_without_reusing_old_native_index(monkeypatch):
    monkeypatch.setattr(settings, "ocr_engine_fingerprint", "regression-fingerprint")
    monkeypatch.setattr(settings, "ocr_executor_url", "")
    payload = office_image("pdf")
    with TestClient(app) as c:
        source = upload(c, "old.pdf", payload, "application/pdf")
        with SessionLocal() as db:
            file = db.get(FileObject, source)
            cache_page(db, file, 0, "cached legacy invoice")
            key = f"derived/{source}/ocr.txt"
            get_storage().put_bytes(key, b"cached legacy invoice", "text/plain")
            file.meta = {
                **file.meta,
                "extract": {**file.meta["extract"], "parser": "zcode-3c5b227-ailab1"},
                "ocr": "done",
                "ocr_text_key": key,
            }
            db.commit()
        from ailab_platform.worker.pipeline import process_file
        from tests.test_file_previews import ApiClient

        class CacheClient(ApiClient):
            def reuse_document(self, fid, scan):
                return self.api.post(f"/internal/files/{fid}/reuse-document", headers=SERVICE, json={"scan": scan}).json()

            def document_cache_pages(self, fid, source_id, after):
                return self.api.get(
                    f"/internal/files/{fid}/document-cache-pages", headers=SERVICE, params={"source_id": source_id, "after": after}
                ).json()

            def publish_cache_pages(self, fid, body):
                response = self.api.post(f"/internal/files/{fid}/document-cache-pages", headers=SERVICE, json=body)
                response.raise_for_status()
                return response.json()

        target = c.post("/v1/files", headers=AUTH, files={"upload": ("new.pdf", payload, "application/pdf")}).json()["id"]
        process_file(target, platform=CacheClient(c))
        with SessionLocal() as db:
            file = db.get(FileObject, target)
            assert file.meta["extract"]["parser"] == "zcode-3c5b227-ailab2"
            assert file.meta["ocr_copied_pages"] == 1
            record = db.scalar(select(OcrPageResult).where(OcrPageResult.file_id == target))
            assert record.text_key.startswith(f"derived/{target}/") and record.seconds == 0
        run = c.post("/v1/runs", headers=AUTH, json={"prompt": "read", "file_ids": [target], "mode": "chat"})
        assert run.status_code == 201
        model_input = c.get(f"/internal/runs/{run.json()['id']}/model-input", headers=SERVICE).json()
        assert model_input["input_files"][0]["text"] == "cached legacy invoice"
    from ailab_platform.internal_skills import load

    assert "不要启动沙箱sleep" in load("platform.pdf.read", "zcode-3c5b227-ailab1")
