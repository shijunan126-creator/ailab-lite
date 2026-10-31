"""Library sandbox rendering reuses public lease lifecycle, including error cleanup."""

import hashlib
import io
import shutil
import tarfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from ailab_platform.broker.documents import DocumentReader
from ailab_platform.broker.render import LibraryRenderer
from ailab_platform.broker.service import LeaseError
from ailab_platform.sandbox.base import LeaseHandle
from tests.test_workspace_previews import _manager


def setup_renderer(tmp_path, content, filename):
    manager, api, storage = _manager(tmp_path)
    api.get_file = lambda fid, **bound: {
        "id": fid,
        "render_owner_id": "user_test",
        "status": "ready",
        "filename": filename,
        "meta": {"preview_job_id": "job", "preview": {"status": "pending"}},
    }
    api.download_file = lambda fid, target, **bound: Path(target).write_bytes(content)
    api.publish_preview_page = lambda *args: None
    execute = manager.provider.exec

    def mock_listener(ref, command, timeout):
        if "warmup.fodt" in command or "render_sandbox_health.py baseline" in command:
            return SimpleNamespace(exit_code=0, stdout='{"tmp_used":0}')
        if "render_sandbox_health.py" in command:
            root = manager.provider.root / ref / "workspace"
            for directory in (root / "inputs", root / ".ailab"):
                for path in directory.iterdir():
                    if path.name in {"uno_convert.py", "render_sandbox_health.py"}:
                        continue
                    if path.is_dir():
                        shutil.rmtree(path)
                    else:
                        path.unlink()
            return SimpleNamespace(exit_code=0, stdout="{}")
        return execute(ref, command.replace("AILAB_LO_STRICT=1 ", ""), timeout)

    manager.provider.exec = mock_listener
    return manager, LibraryRenderer(manager), storage


def published(manager, result):
    lease_id = result["render_lease_id"]
    lease = manager.get(lease_id)
    assert manager.finish_render(
        LeaseHandle(lease_id, lease["provider_ref"]), run_id="preview:job", generation=result["render_generation"], success=True
    )


@pytest.mark.skipif(not all(shutil.which(x) for x in ["soffice", "pdftoppm", "pdfinfo", "timeout"]), reason="Office tools unavailable")
def test_docx_real_renderer_and_release(tmp_path):
    doc = Document()
    doc.add_paragraph("Library preview contract")
    data = io.BytesIO()
    doc.save(data)
    manager, renderer, storage = setup_renderer(tmp_path, data.getvalue(), "hello.docx")
    result = renderer.render("file", "job")
    assert result["preview"]["status"] == "ready" and result["preview"]["page_count"] == 1
    assert len(result["preview"]["pages"]) == 1
    key = next(iter(result["assets"].values()))["key"]
    assert storage.get_bytes(key).startswith(b"\x89PNG\r\n\x1a\n")
    published(manager, result)
    manager.release(result["render_lease_id"])
    assert manager.active_leases() == []


def test_failed_renderer_releases_and_job_id_is_bound(tmp_path):
    manager, renderer, _ = setup_renderer(tmp_path, b"broken", "bad.pdf")
    with pytest.raises(LeaseError, match="pending"):
        renderer.render("file", "wrong")
    with pytest.raises(LeaseError, match="render_failed"):
        renderer.render("file", "job")
    assert manager.active_leases() == []


@pytest.mark.skipif(not all(shutil.which(x) for x in ["node", "pdftoppm", "pdfinfo", "timeout"]), reason="reader tools unavailable")
def test_parse_reuses_same_user_sandbox_and_job_binding(tmp_path):
    source = Path(__file__).resolve().parents[2] / "vendor/zcode-office-runtime/fixtures/paper.pdf"
    manager, renderer, storage = setup_renderer(tmp_path, source.read_bytes(), "paper.pdf")
    preview = renderer.render("file", "job")
    published(manager, preview)
    original = manager.platform.get_file
    manager.platform.get_file = lambda fid, **binding: {**original(fid, **binding), "size": source.stat().st_size}
    reads = []
    download = manager.platform.download_file
    manager.platform.download_file = lambda fid, target, **binding: reads.append(binding) or download(fid, target, **binding)
    parsed = DocumentReader(manager).read("file", "document-job")
    assert parsed["render_lease_id"] == preview["render_lease_id"]
    assert parsed["lease_source"] == "warm" and storage.exists(parsed["archive_key"])
    assert reads[0]["job"] == "document" and reads[0]["preview_job"] == "document-job"
    lease = manager.get(parsed["render_lease_id"])
    assert manager.finish_render(
        LeaseHandle(parsed["render_lease_id"], lease["provider_ref"]),
        run_id="document:document-job",
        generation=parsed["render_generation"],
        success=True,
    )
    manager.release(parsed["render_lease_id"])


@pytest.mark.skipif(not all(shutil.which(x) for x in ["pdftoppm", "pdfinfo", "timeout", "tar"]), reason="PDF tools unavailable")
def test_first_page_is_published_before_remaining_render(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    data = io.BytesIO()
    writer.write(data)
    manager, renderer, _ = setup_renderer(tmp_path, data.getvalue(), "pages.pdf")
    events = []
    execute = manager.exec
    manager.exec = lambda lease, cmd, timeout: (
        events.append("first" if "/first " in cmd else "second" if "/page-2 " in cmd else "other") or execute(lease, cmd, timeout)
    )
    manager.platform.publish_preview_page = lambda fid, job, name, count, asset: events.append(f"publish:{name}")
    result = renderer.render("file", "job")
    assert events.index("first") < events.index("publish:p-1.png") < events.index("second") < events.index("publish:p-2.png")
    assert events.count("publish:p-1.png") == 1
    published(manager, result)
    manager.release(result["render_lease_id"])


def test_waiting_page_is_rendered_before_other_remaining_pages(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(4):
        writer.add_blank_page(width=72, height=72)
    data = io.BytesIO()
    writer.write(data)
    manager, renderer, _ = setup_renderer(tmp_path, data.getvalue(), "pages.pdf")
    requested, published_pages = [], []
    original = manager.platform.get_file
    manager.platform.get_file = lambda fid, **binding: {
        **original(fid, **binding),
        "meta": {**original(fid, **binding)["meta"], "preview_requested_pages": requested},
    }

    def publish(fid, job, name, count, asset):
        published_pages.append(int(name[2:-4]))
        if len(published_pages) == 1:
            requested.append(4)

    manager.platform.publish_preview_page = publish
    result = renderer.render("file", "job")
    assert published_pages == [1, 4, 2, 3]
    assert [int(name[2:-4]) for name in result["preview"]["pages"]] == [1, 2, 3, 4]
    published(manager, result)
    manager.release(result["render_lease_id"])


@pytest.mark.skipif(not all(shutil.which(x) for x in ["pdftoppm", "pdfinfo", "timeout"]), reason="PDF tools unavailable")
def test_pdf_caps_at_40_but_reports_total_page_count(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(57):
        writer.add_blank_page(width=72, height=72)
    data = io.BytesIO()
    writer.write(data)
    manager, renderer, _ = setup_renderer(tmp_path, data.getvalue(), "57pages.pdf")
    result = renderer.render("file", "job")
    assert result["preview"]["page_count"] == 57 and len(result["preview"]["pages"]) == 40
    published(manager, result)
    manager.release(result["render_lease_id"])
    assert manager.active_leases() == []


@pytest.mark.skipif(not all(shutil.which(x) for x in ["pdftoppm", "pdfinfo", "timeout", "tar"]), reason="PDF tools unavailable")
def test_render_exports_and_publishes_each_page(tmp_path, caplog):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=72, height=72)
    data = io.BytesIO()
    writer.write(data)
    manager, renderer, storage = setup_renderer(tmp_path, data.getvalue(), "pages.pdf")
    downloads = []
    original = manager._download

    def record_download(lease, path, local, limit):
        downloads.append(path)
        return original(lease, path, local, limit)

    manager._download = record_download
    with caplog.at_level("INFO", logger="ailab.broker.render"):
        result = renderer.render("file", "job")
    assert len(downloads) == 3 and downloads[0].startswith("/workspace/.ailab/first/p-")
    assert all(f"/page-{i}/" in downloads[i - 1] for i in (2, 3))
    assert len(result["assets"]) == 3
    assert "convert_seconds" in caplog.text and "raster_seconds" in caplog.text
    assert all(storage.get_bytes(asset["key"]).startswith(b"\x89PNG\r\n\x1a\n") for asset in result["assets"].values())
    published(manager, result)
    manager.release(result["render_lease_id"])
    assert manager.active_leases() == []


@pytest.mark.skipif(not all(shutil.which(x) for x in ["pdftoppm", "pdfinfo", "timeout", "tar"]), reason="PDF tools unavailable")
def test_render_rejects_unexpected_archive_members(tmp_path):
    manager, renderer, _ = setup_renderer(tmp_path, b"a,b\n1,2\n", "sheet.csv")
    original = manager._download

    def tamper(lease, path, local, limit):
        original(lease, path, local, limit)
        with tarfile.open(local, "a") as bundle:
            entry = tarfile.TarInfo("../escape")
            entry.size = 1
            bundle.addfile(entry, io.BytesIO(b"x"))

    manager._download = tamper
    with pytest.raises(LeaseError, match="invalid_render_output"):
        renderer.render("file", "job")
    assert manager.active_leases() == []


@pytest.mark.skipif(
    not all(shutil.which(x) for x in ["soffice", "pdftoppm", "pdfinfo", "timeout", "tar"]), reason="Office tools unavailable"
)
def test_visual_page_requests_reuse_private_pdf_and_keep_requested_order(tmp_path):
    document = Document()
    document.add_paragraph("First")
    document.add_page_break()
    document.add_paragraph("Second")
    data = io.BytesIO()
    document.save(data)
    content = data.getvalue()
    manager, renderer, storage = setup_renderer(tmp_path, content, "hello.docx")
    original = manager.platform.get_file
    manager.platform.get_file = lambda fid, **bound: {
        **original(fid, **bound),
        "sha256": hashlib.sha256(content).hexdigest(),
        "meta": {"preview_job_id": "job", "preview": {"status": "pending"}, "ocr_render_job_id": "job", "ocr_render_pages": [2, 1]},
    }
    commands = []
    execute = manager.exec
    manager.exec = lambda lease, cmd, timeout: commands.append(cmd) or execute(lease, cmd, timeout)
    first = renderer.render("file", "job")
    published(manager, first)
    second = renderer.render("file", "job")
    published(manager, second)
    other = renderer.render("other-file", "job")
    published(manager, other)
    assert [int(name[2:-4]) for name in first["preview"]["pages"]] == [2, 1]
    assert second["preview"] == first["preview"]
    assert [first["assets"][name]["sha256"] for name in first["preview"]["pages"]] == [
        second["assets"][name]["sha256"] for name in second["preview"]["pages"]
    ]
    assert [first["assets"][name]["sha256"] for name in first["preview"]["pages"]] == [
        other["assets"][name]["sha256"] for name in other["preview"]["pages"]
    ]
    render_commands = [cmd for cmd in commands if "render_preview.py" in cmd]
    assert other["preview"] == first["preview"]
    assert len(render_commands) == 6
    assert "/workspace/inputs/input.docx" in render_commands[0]
    assert "/workspace/inputs/cached.pdf" in render_commands[2]
    assert "/workspace/inputs/input.docx" in render_commands[4]
    assert storage.list_prefix("derived/file/render-cache/")


def test_render_uses_the_reserved_render_class_and_reads_bound_to_its_job(tmp_path):
    """A render is a system lease of the render profile; every file read names the preview job."""
    manager, renderer, _ = setup_renderer(tmp_path, b"broken", "bad.pdf")
    reads, acquired = [], []
    original_get = manager.platform.download_file
    manager.platform.download_file = lambda fid, target, **bound: reads.append(bound) or original_get(fid, target, **bound)
    original_acquire = manager.acquire
    manager.acquire = lambda spec, principal, **kw: acquired.append((spec, principal)) or original_acquire(spec, principal, **kw)
    with pytest.raises(LeaseError, match="render_failed"):
        renderer.render("file", "job")
    spec, principal = acquired[0]
    assert spec.profile == "render" and principal.system and spec.max_lifetime <= 120
    assert reads and all(r.get("job") == "preview" and r.get("preview_job") == "job" for r in reads)
    assert manager.active_leases() == []


@pytest.mark.skipif(not all(shutil.which(x) for x in ["soffice", "pdftoppm", "pdfinfo", "timeout"]), reason="Office tools unavailable")
@pytest.mark.parametrize("potx", [False, True])
def test_template_real_sandbox_preview_includes_cached_first_page_thumbnail(tmp_path, potx):
    from PIL import Image

    from tests.test_templates import pptx

    filename = "template.potx" if potx else "template.pptx"
    manager, renderer, storage = setup_renderer(tmp_path, pptx(slides=2, potx=potx), filename)
    original = manager.platform.get_file
    manager.platform.get_file = lambda fid, **bound: {
        **original(fid, **bound),
        "meta": {"template_kind": "pptx", "preview_job_id": "job", "preview": {"status": "pending", "kind": "pages"}},
    }
    result = renderer.render("file", "job")
    assert len(result["preview"]["pages"]) == 2
    assert set(result["assets"]) == {*result["preview"]["pages"], "template-thumbnail.jpg"}
    asset = result["assets"]["template-thumbnail.jpg"]
    assert asset["key"] == "derived/file/preview/template-thumbnail.jpg"
    with Image.open(io.BytesIO(storage.get_bytes(asset["key"]))) as image:
        assert image.format == "JPEG" and image.width == 480
    published(manager, result)
    manager.release(result["render_lease_id"])
    assert manager.active_leases() == []
