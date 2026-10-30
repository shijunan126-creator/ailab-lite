"""Deterministic fixtures and checker for externally generated golden outputs; no paid calls."""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]


def check_outputs(directory: Path):
    deck = Presentation(directory / "deck-basic.pptx")
    assert len(deck.slides) == 2, "slide count"
    texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                assert shape.left >= 0 and shape.top >= 0, "negative text shape origin"
                assert shape.left + shape.width <= deck.slide_width + 10, "text shape outside page"
                assert shape.top + shape.height <= deck.slide_height + 10, "text shape outside page"
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        assert run.font.name, "missing explicit font"
    assert any("Repository golden" in text for text in texts), "missing title"
    assert any("Checks" in text for text in texts), "missing second-slide text"
    book = load_workbook(directory / "workbook-basic.xlsx")
    assert "Data" in book.sheetnames, "missing sheet"
    values = [cell.value for row in book["Data"] for cell in row]
    assert 2 in values and 3 in values, "input values missing"
    assert any(cell.data_type == "f" and cell.value.startswith("=SUM(") for row in book["Data"] for cell in row), "missing SUM formula"
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "e", f"Excel error: {cell.coordinate}"
    # Cached errors are independently checked; formula recalculation requires the real Office runner.
    cached = load_workbook(directory / "workbook-basic.xlsx", data_only=True)
    assert not any(cell.data_type == "e" for sheet in cached for row in sheet for cell in row), "cached formula error"
    return {"samples": 2, "scope": "repository fixtures; structural assertions only; business and rendered overflow unverified"}


def self_test(directory: Path):
    for name, script, suffix in [
        ("deck-basic", "platform.pptx/scripts/build.py", "pptx"),
        ("workbook-basic", "platform.xlsx/scripts/build_xlsx.py", "xlsx"),
    ]:
        subprocess.run(
            [
                sys.executable,
                str(ROOT / "images/sbx-office/skills" / script),
                str(HERE / name / "input.json"),
                str(directory / f"{name}.{suffix}"),
            ],
            check=True,
            timeout=60,
        )
    return check_outputs(directory)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--outputs", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    if args.self_test:
        with tempfile.TemporaryDirectory(prefix="ailab-golden-") as tmp:
            result = self_test(Path(tmp))
    elif args.outputs:
        result = check_outputs(args.outputs)
    else:
        parser.error("choose --self-test or --outputs")
    report = json.dumps(result, ensure_ascii=False, indent=2)
    print(report)
    if args.report:
        args.report.write_text(report + "\n")
