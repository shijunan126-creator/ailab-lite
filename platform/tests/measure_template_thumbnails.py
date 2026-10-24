"""Opt-in real Office cost probe; caller imports tests.conftest for isolation first."""

import json
import os
import statistics
import tempfile
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor

from ailab_platform.config import settings
from ailab_platform.storage import LocalStorage
from ailab_platform.worker.isolate import IsolationError, execute
from ailab_platform.worker.pipeline import _template_thumbnail


def main():
    with tempfile.TemporaryDirectory(prefix="ailab-template-render-measure-") as tmp:
        root = Path(tmp)
        source = root / "template.pptx"
        prs = Presentation()
        slides = int(os.environ.get("AILAB_MEASURE_SLIDES", "1"))
        for index in range(slides):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = f"Template sample {index + 1}"
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string("123456" if index == 0 else "AA3322")
        prs.save(source)
        result = {"input_bytes": source.stat().st_size, "slides": slides, "profiles": {}}
        for memory_mb in (192, 256):
            samples = []
            for index in range(3):
                started = time.perf_counter()
                try:
                    value = execute(
                        "template_thumbnail",
                        source,
                        root / f"{memory_mb}-{index}",
                        memory_mb=memory_mb,
                        timeout=settings.parse_timeout_seconds * 0.85,
                        watch_process_tree=True,
                        measure=True,
                    )
                    sample = {"status": value["thumbnail"], **value["measurement"]}
                except IsolationError as exc:
                    sample = {"status": str(exc)}
                samples.append({**sample, "wall_seconds": time.perf_counter() - started})
            successes = [sample for sample in samples if sample["status"] == "rendered"]
            result["profiles"][memory_mb] = {
                "samples": samples,
                "wall_seconds_median": statistics.median(sample["wall_seconds"] for sample in samples),
                "cpu_seconds_median": statistics.median(sample["cpu_seconds"] for sample in successes) if successes else None,
                "peak_process_tree_rss_bytes_max": max(sample["peak_process_tree_rss_bytes"] for sample in successes)
                if successes
                else None,
            }
        started = time.perf_counter()
        meta = _template_thumbnail("fil_render_probe", source, LocalStorage(root / "storage"), root / "fallback")
        result["configured_pipeline"] = {"memory_mb": settings.parse_memory_mb, "wall_seconds": time.perf_counter() - started, **meta}
        print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
