"""Project/brand API contracts, bounded analysis, immutable bootstrap and office output."""

import asyncio
import importlib.util
import io
import json
import uuid
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import text

from ailab_platform import accounts
from ailab_platform.agent.loop import AgentLoop
from ailab_platform.api.main import app
from ailab_platform.api.routes_internal import _bound_file
from ailab_platform.db import SessionLocal
from ailab_platform.llm import LLMResponse
from ailab_platform.models import FileObject, PlatformSetting, Run, UserRole
from ailab_platform.quota import library_bytes
from ailab_platform.service_auth import ServicePrincipal
from ailab_platform.worker.pipeline import process_file
from tests.run_claims import attempt_headers
from tests.test_agent_loop import FakeBroker, FakePlatform, ScriptedLLM
from tests.test_broker_lease import _ApiStub, _manager_with
from tests.test_migrations import engine as engine
from tests.test_projects import artifact

SERVICE = {"X-Service-Token": "test-service"}


@pytest.fixture
def api():
    with TestClient(app) as client:
        with SessionLocal() as db:
            user = accounts.create_user(db, "brand-" + uuid.uuid4().hex[:12], "password-for-testing")
            user.role_rows = [UserRole(role="user"), UserRole(role="platform_admin")]
            token, _ = accounts.issue_session(db, user)
            db.commit()
        client.headers.update({"Authorization": "Bearer " + token})
        # Tests share a database but must not share platform settings.
        with SessionLocal() as db:
            row = db.get(PlatformSetting, "brand")
            if row:
                db.delete(row)
            db.commit()
        yield client


def png():
    out = io.BytesIO()
    Image.new("RGB", (20, 10), "red").save(out, "PNG")
    return out.getvalue()


def upload_brand(api, kind="logo", data=None, name="logo.png"):
    response = api.post(f"/v1/admin/brand/assets/{kind}", files={"file": (name, png() if data is None else data)})
    assert response.status_code == 201, response.text
    return response.json()


class WorkerAPI:
    def __init__(self, api):
        self.api = api

    def get_file(self, file_id, **kwargs):
        response = self.api.get(f"/internal/files/{file_id}", headers=SERVICE)
        assert response.status_code == 200, response.text
        return response.json()

    def download_file(self, file_id, target, max_bytes, **kwargs):
        response = self.api.get(f"/internal/files/{file_id}/content", headers=SERVICE)
        assert response.status_code == 200, response.text
        assert len(response.content) <= max_bytes
        target.write_bytes(response.content)

    def patch_file(self, file_id, **kwargs):
        response = self.api.patch(f"/internal/files/{file_id}", headers=SERVICE, json=kwargs)
        assert response.status_code == 200, response.text
        return response.json()


def start(api, project_id=None, mode="work"):
    conversation = api.post("/v1/conversations", json={"project_id": project_id}).json()
    response = api.post(f"/v1/conversations/{conversation['id']}/messages", json={"prompt": "test", "mode": mode})
    assert response.status_code == 201, response.text
    return conversation, response.json()


def test_project_instructions_patch_and_both_bootstraps(api):
    project = api.post("/v1/projects", json={"name": "Instructions"}).json()
    path = f"/v1/projects/{project['id']}"
    assert project["instructions"] == ""
    assert api.patch(path, json={"instructions": "x" * 4001}).status_code == 422
    assert api.patch(path, json={"name": None}).status_code == 422
    response = api.patch(path, json={"instructions": "Keep the project scope", "name": "Renamed"})
    assert response.json()["instructions"] == "Keep the project scope"
    _, run = start(api, project["id"], mode="chat")
    for suffix in ("", "/model-input"):
        out = api.get(f"/internal/runs/{run['id']}{suffix}", headers=SERVICE).json()
        assert out["project_instructions"] == "Keep the project scope"
        assert out["brand"] == {}
    assert api.patch(path, json={"instructions": ""}).json()["instructions"] == ""


def test_project_instructions_and_brand_are_frozen_in_checkpoint():
    class Platform(FakePlatform):
        instruction = "first instruction"

        async def get_run(self, run_id):
            return {
                **await super().get_run(run_id),
                "project_instructions": self.instruction,
                "brand": {"version": 3, "assets": {}},
                "memory": {"entries": [], "enabled": True},
            }

    platform = Platform()
    llm = ScriptedLLM([LLMResponse(content="done")])
    asyncio.run(AgentLoop(platform=platform, broker=FakeBroker(), llm=llm).process("run_1", attempt_id="att_1"))
    frozen = platform.checkpoint["state"]["messages"][1]["content"]
    assert "first instruction" in frozen and '"version": 3' in frozen
    assert "revision=" not in frozen
    platform.checkpoint["state"]["finished"] = False
    platform.instruction = "changed instruction"
    llm = ScriptedLLM([LLMResponse(content="done again")])
    asyncio.run(AgentLoop(platform=platform, broker=FakeBroker(), llm=llm).process("run_1", attempt_id="att_2"))
    assert llm.seen_messages[0][1]["content"] == frozen


def test_brand_settings_validation_and_permissions(api):
    assert api.get("/v1/admin/brand").json()["enabled"] is False
    assert api.patch("/v1/admin/brand", json={"colors": {"primary": "red", "accent": "#FF0000"}}).status_code == 422
    assert api.patch("/v1/admin/brand", json={"font": "x" * 61}).status_code == 422
    assert api.patch("/v1/admin/brand", json={"enabled": None}).status_code == 422
    assert api.post("/v1/admin/brand/assets/logo", files={"file": ("logo.svg", b"<svg/>")}).status_code == 422
    assert api.post("/v1/admin/brand/assets/logo", files={"file": ("logo.png", b"x" * (2 * 1024 * 1024 + 1))}).status_code == 413
    result = api.patch(
        "/v1/admin/brand", json={"enabled": True, "name": "Example", "font": "Arial", "colors": {"primary": "#123456", "accent": "#ABCDEF"}}
    ).json()
    assert result["font"] == "Arial" and result["version"] > 0
    with SessionLocal() as db:
        user = accounts.create_user(db, "reader-" + uuid.uuid4().hex[:12], "password-for-testing")
        token, _ = accounts.issue_session(db, user)
        db.commit()
    for method, suffix, kwargs in [
        ("GET", "", {}),
        ("PATCH", "", {"json": {"enabled": False}}),
        ("POST", "/assets/logo", {"files": {"file": ("logo.png", png())}}),
        ("DELETE", "/assets/logo", {}),
        ("GET", "/assets/logo/content", {}),
    ]:
        assert api.request(method, "/v1/admin/brand" + suffix, headers={"Authorization": "Bearer " + token}, **kwargs).status_code == 403


def test_brand_scan_analysis_no_library_quota_and_immutable_replacement(api):
    api.patch("/v1/admin/brand", json={"enabled": True})
    first = upload_brand(api)
    assert first["status"] == "pending"
    _, pending_run = start(api)
    empty = api.get(f"/internal/runs/{pending_run['id']}", headers=SERVICE).json()["brand"]
    assert empty["assets"] == {}
    result = process_file(first["file_id"], WorkerAPI(api))
    assert result["brand"]["width"] == 20
    settings = api.get("/v1/admin/brand").json()
    assert settings["logo_file_id"] == first["file_id"] and settings["assets"]["logo"]["status"] == "ready"
    assert api.get("/v1/admin/brand/assets/logo/content").content == png()
    assert api.get("/v1/admin/brand/assets/logo/thumbnail").status_code == 200
    assert api.get(f"/v1/files/{first['file_id']}/content").status_code == 404
    assert all(f["id"] != first["file_id"] for f in api.get("/v1/files").json())
    assert api.post("/v1/files/batch", json={"action": "trash", "ids": [first["file_id"]]}).status_code == 404
    with SessionLocal() as db:
        file = db.get(FileObject, first["file_id"])
        assert file.purpose == "brand" and library_bytes(db, file.user_id) == 0
    _, run = start(api)
    snapshot = api.get(f"/internal/runs/{run['id']}/model-input", headers=SERVICE).json()["brand"]
    second = upload_brand(api)
    assert api.get("/v1/admin/brand").json()["logo_file_id"] == first["file_id"]
    process_file(second["file_id"], WorkerAPI(api))
    assert api.get("/v1/admin/brand").json()["logo_file_id"] == second["file_id"]
    assert api.get(f"/internal/runs/{run['id']}", headers=SERVICE).json()["brand"] == snapshot
    assert api.get(f"/internal/runs/{pending_run['id']}", headers=SERVICE).json()["brand"] == empty
    with SessionLocal() as db:
        assert db.get(Run, run["id"]).runtime_meta["brand_version"] == snapshot["version"]
        broker = ServicePrincipal("broker", frozenset({"files:read_for_run"}))
        agent = ServicePrincipal("agent", frozenset({"files:read_for_run"}))
        assert _bound_file(db, first["file_id"], broker, run["id"], "brand").id == first["file_id"]
        from fastapi import HTTPException

        for principal, job in [(broker, "processing"), (agent, "brand")]:
            with pytest.raises(HTTPException):
                _bound_file(db, first["file_id"], principal, run["id"], job)
    assert api.delete("/v1/admin/brand/assets/logo").status_code == 204
    _, new = start(api)
    assert api.get(f"/internal/runs/{new['id']}", headers=SERVICE).json()["brand"]["assets"] == {}


@pytest.mark.parametrize("kind", ["pptx", "docx"])
def test_templates_analyzed_once_in_limited_child(api, kind, tmp_path):
    source = tmp_path / ("template." + kind)
    if kind == "pptx":
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.save(source)
    else:
        Document().save(source)
    if kind == "pptx":
        response = api.post("/v1/templates", files={"upload": (source.name, source.read_bytes())})
        assert response.status_code == 201, response.text
        with SessionLocal() as db:
            from ailab_platform.models import Template

            asset = {"file_id": db.get(Template, response.json()["id"]).file_id}
    else:
        asset = upload_brand(api, kind, source.read_bytes(), source.name)
    result = process_file(asset["file_id"], WorkerAPI(api))
    if kind == "pptx":
        assert result["template"]["slide_width"] == Inches(10)
        assert result["template"]["master_count"] == 1 and result["template"]["themes"][0]["colors"]
    else:
        assert "Normal" in result["brand"]["styles"] and len(result["brand"]["styles"]) <= 50
    assert process_file(asset["file_id"], WorkerAPI(api))["skipped"] == "ready"


def test_corrupt_brand_rejected_without_replacing_ready(api):
    first = upload_brand(api)
    process_file(first["file_id"], WorkerAPI(api))
    second = upload_brand(api, "docx", b"not a zip", "broken.docx")
    assert process_file(second["file_id"], WorkerAPI(api))["rejected"] == "corrupt"
    assert api.get("/v1/admin/brand").json()["logo_file_id"] == first["file_id"]


def test_move_conversation_artifact_lineages_in_one_transaction(api):
    project = api.post("/v1/projects", json={"name": "Destination"}).json()
    conv, run = start(api)
    api.post(f"/internal/runs/{run['id']}/status", headers=attempt_headers(api, run["id"]), json={"status": "running"})
    out = artifact(api, run["id"])
    assert api.patch(f"/v1/conversations/{conv['id']}", json={"move_files": True}).status_code == 422
    bad = api.patch(f"/v1/conversations/{conv['id']}", json={"project_id": "missing", "move_files": True})
    assert bad.status_code == 404
    with SessionLocal() as db:
        assert db.get(FileObject, out["file_id"]).project_id is None
    response = api.patch(f"/v1/conversations/{conv['id']}", json={"project_id": project["id"], "move_files": True})
    assert response.json()["moved_files"] == 1
    with SessionLocal() as db:
        assert db.get(FileObject, out["file_id"]).project_id == project["id"]
    response = api.patch(f"/v1/conversations/{conv['id']}", json={"project_id": None})
    assert response.json()["moved_files"] is None
    with SessionLocal() as db:
        assert db.get(FileObject, out["file_id"]).project_id == project["id"]


def test_brand_broker_injects_only_frozen_assets_with_limits(monkeypatch):
    brand = {"version": 2, "colors": {"primary": "#123456"}, "assets": {"logo": "logo"}}

    class API(_ApiStub):
        def get_run(self, run_id):
            return {**super().get_run(run_id), "brand": brand, "options": {"mode": "work"}}

    api = API({"logo": {"status": "ready", "size": 5, "bytes": b"image"}}, [])
    manager, provider, _ = _manager_with(api)
    calls = []
    original = api.download_file

    def download(file_id, target, **binding):
        calls.append(binding)
        return original(file_id, target, **binding)

    monkeypatch.setattr(api, "download_file", download)
    lease = manager.create("brand-run")
    assert provider.get_file(lease["sandbox_id"], "/workspace/.brand/logo.png") == b"image"
    assert json.loads(provider.get_file(lease["sandbox_id"], "/workspace/.brand/brand.json"))["version"] == 2
    assert calls == [{"run_id": "brand-run", "job": "brand", "max_bytes": 2 * 1024 * 1024}]
    manager.release(lease["lease_id"])
    api.files["logo"]["size"] = 2 * 1024 * 1024 + 1
    from ailab_platform.broker.service import LeaseError

    with pytest.raises(LeaseError):
        manager.create("too-large-brand")
    api.files["logo"].update(size=5, bytes=b"x" * (2 * 1024 * 1024 + 1))
    with pytest.raises(LeaseError):
        manager.create("lying-brand-size")


def test_generator_applies_brand_and_template_and_opt_out(tmp_path, monkeypatch):
    path = Path(__file__).resolve().parents[2] / "images/sbx-office/skills/platform.pptx/scripts/build.py"
    local_script = tmp_path / "build.py"
    local_script.write_bytes(path.read_bytes())
    spec = importlib.util.spec_from_file_location("brand_ppt_builder", local_script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(module, "BRAND_DIR", tmp_path)
    (tmp_path / "brand.json").write_text(json.dumps({"colors": {"primary": "#123456", "accent": "#ABCDEF"}, "font": "Arial"}))
    (tmp_path / "logo.png").write_bytes(png())
    template = Presentation()
    template.slide_width, template.slide_height = Inches(10), Inches(7.5)
    template.slides.add_slide(template.slide_layouts[6])
    template.save(tmp_path / "template.pptx")
    deck = module.Deck({"slides": [{"layout": "cover", "title": "Brand"}, {"layout": "bullets", "title": "Content", "bullets": ["Body"]}]})
    assert deck.t["primary"] == "123456" and module.FONT == "Arial"
    deck.build(str(tmp_path / "out.pptx"))
    result = Presentation(tmp_path / "out.pptx")
    assert len(result.slides) == 2 and result.slide_width == Inches(10)
    assert len(result.slide_masters) == len(template.slide_masters)
    assert all(any(shape.shape_type == 13 for shape in slide.shapes) for slide in result.slides)
    fonts = [
        r.font.name
        for slide in result.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
        for r in p.runs
    ]
    assert fonts and all(font == "Arial" for font in fonts)
    plain = module.Deck({"brand": False, "slides": []})
    assert plain.prs.slide_width == module.W and module.FONT == "微软雅黑" and not plain.brand


def test_brand_new_queries_use_indexes():
    # Run in the active suite database, SQLite and PG; primary-key bootstrap/settings lookups need no new index.
    with SessionLocal() as db:
        prefix = "EXPLAIN QUERY PLAN " if db.bind.dialect.name == "sqlite" else "EXPLAIN "
        if db.bind.dialect.name == "postgresql":
            db.execute(text("SET LOCAL enable_seqscan = off"))
        for sql, expected in [
            ("SELECT value FROM platform_settings WHERE key='brand'", "platform_settings"),
            ("SELECT instructions FROM projects WHERE id='missing'", "projects"),
            (
                "SELECT f.id FROM files f WHERE f.is_head=1 AND f.lineage_id IN "
                "(SELECT f2.lineage_id FROM runs r JOIN artifacts a ON a.run_id=r.id JOIN files f2 ON f2.id=a.file_id "
                "WHERE r.conversation_id='missing' AND r.user_id='missing' AND f2.user_id='missing' "
                "AND f2.purpose='library' AND f2.deleted_at IS NULL AND f2.purged_at IS NULL)",
                "index",
            ),
        ]:
            plan = str(db.execute(text(prefix + sql)).all()).lower()
            assert expected in plan and ("index" in plan or "primary key" in plan), plan


def test_project_schedule_bootstrap_uses_same_instructions(api):
    from tests.test_schedules_citations import body

    project = api.post("/v1/projects", json={"name": "Scheduled project"}).json()
    api.patch(f"/v1/projects/{project['id']}", json={"instructions": "Scheduled instructions"})
    task = api.post("/v1/schedules", json=body(project_id=project["id"], enabled=False, mode="work")).json()
    response = api.post(f"/v1/schedules/{task['id']}/run")
    assert response.status_code == 201, response.text
    run_id = response.json()["run_id"]
    assert api.get(f"/internal/runs/{run_id}/model-input", headers=SERVICE).json()["project_instructions"] == "Scheduled instructions"


def test_project_instructions_migration_roundtrip(engine):
    from alembic import command
    from sqlalchemy import inspect

    from ailab_platform.migrate import alembic_config

    cfg = alembic_config()
    for revision, direction in [("0024", "upgrade"), ("0025", "upgrade"), ("0024", "downgrade"), ("0025", "upgrade")]:
        with engine.begin() as conn:
            cfg.attributes["connection"] = conn
            getattr(command, direction)(cfg, revision)
        columns = {col["name"] for col in inspect(engine).get_columns("projects")}
        assert ("instructions" in columns) == (revision == "0025")


def test_infected_brand_is_not_analyzed_and_has_visible_reason(api):
    asset = upload_brand(api, data=b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE")
    assert asset["analysis"] is None
    result = process_file(asset["file_id"], WorkerAPI(api))
    assert result["rejected"]["verdict"] == "infected"
    settings = api.get("/v1/admin/brand").json()
    assert settings["logo_file_id"] is None
    assert settings["assets"]["logo"]["analysis"]["error"] == "infected"
