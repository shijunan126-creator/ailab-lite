"""Opt-in, disposable five-service load on the dedicated Docker verification VM.

Requires ailab-final-backend:verify and ailab-final-portal:verify built from the
same source, plus redis:7-alpine. No paid models, production data or credentials.
The load generator stays on the host; service CPU/memory comes from cgroups.
"""

import argparse
import concurrent.futures
import hashlib
import json
import os
import socket
import statistics
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

import httpx
from pptx import Presentation


def free_port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def measure(context, chats):
    prefix = "ailab-mixed-" + uuid.uuid4().hex[:8]
    network, volume = prefix + "-net", prefix + "-data"
    names = []

    def docker(*args):
        return subprocess.check_output(["docker", "--context", context, *args], text=True, stderr=subprocess.PIPE, timeout=180)

    def wait(url):
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            try:
                if httpx.get(url, timeout=2).status_code == 200:
                    return
            except httpx.HTTPError:
                pass
            time.sleep(0.3)
        raise RuntimeError("verification service did not become ready")

    def counters(name):
        # memory.peak includes page cache and parser children. VmRSS is the
        # main process at the sample time, not a claim about peak process RSS.
        text = docker(
            "exec",
            prefix + "-" + name,
            "sh",
            "-c",
            "cat /sys/fs/cgroup/cpu.stat; "
            "cat /sys/fs/cgroup/memory.peak; cat /sys/fs/cgroup/memory.events; "
            "awk '/VmRSS/{print $2}' /proc/1/status",
        )
        lines = text.splitlines()
        cpu = int(next(line.split()[1] for line in lines if line.startswith("usage_usec ")))
        peak = int(next(line for line in lines if line.isdigit()))
        return {
            "cpu_seconds": cpu / 1e6,
            "memory_peak_mib": peak / 1024**2,
            "rss_end_mib": int(lines[-1]) / 1024,
            "oom_kill": int(next(line.split()[1] for line in lines if line.startswith("oom_kill "))),
        }

    api_port, portal_port = free_port(), free_port()
    endpoint = f"http://127.0.0.1:{portal_port}"
    tokens = {name: uuid.uuid4().hex for name in ("api", "agent", "broker", "worker", "ops")}
    password = uuid.uuid4().hex
    try:
        docker("network", "create", network)
        docker("volume", "create", volume)
        with tempfile.TemporaryDirectory(prefix="ailab-mixed-fixture-") as directory:
            root = Path(directory)
            common = {
                "AILAB_ENV": "test",
                "AILAB_ENV_FILE": "/no-such.env",
                "AILAB_DATA_DIR": "/data",
                "AILAB_STORAGE_ROOT": "/data/storage",
                "AILAB_DATABASE_URL": "sqlite:////data/mixed.db",
                "AILAB_QUEUE_BACKEND": "redis",
                "AILAB_REDIS_URL": f"redis://{prefix}-redis:6379/0",
                "AILAB_API_BASE_URL": f"http://{prefix}-api:8080",
                "AILAB_BROKER_BASE_URL": f"http://{prefix}-broker:8080",
                "AILAB_MODELS_FILE": "/etc/ailab/models.toml",
                "AILAB_SANDBOX_PROVIDER": "mock",
                "AILAB_IMAGE_SEARCH_PROVIDER": "mock",
                "AILAB_MAX_ACTIVE_RUNS_PER_USER": "1000",
                "AILAB_DEV_TOKEN": uuid.uuid4().hex,
                "AILAB_ADMIN_USERNAME": "load-probe",
                "AILAB_ADMIN_PASSWORD": password,
                "AILAB_SERVICE_TOKEN": "",
                "AILAB_SERVICE_CREDENTIALS": json.dumps(
                    {name: "sha256:" + hashlib.sha256(token.encode()).hexdigest() for name, token in tokens.items()}
                ),
            }
            limits = {"api": ("0.5", "384m"), "broker": ("0.5", "384m"), "agent": ("0.3", "256m"), "worker": ("0.3", "384m")}
            names.append(prefix + "-redis")
            docker(
                "run",
                "-d",
                "--name",
                names[-1],
                "--network",
                network,
                "--memory",
                "64m",
                "--cpus",
                "0.1",
                "redis:7-alpine",
                "redis-server",
                "--save",
                "",
                "--appendonly",
                "no",
            )
            for role, (cpu, memory) in limits.items():
                envfile = root / (role + ".env")
                envfile.write_text(
                    "\n".join(
                        f"{key}={value}"
                        for key, value in {**common, "AILAB_SERVICE_NAME": role, "AILAB_SERVICE_CREDENTIAL": tokens[role]}.items()
                    )
                    + "\n"
                )
                names.append(prefix + "-" + role)
                publish = ["-p", f"127.0.0.1:{api_port}:8080"] if role == "api" else []
                docker(
                    "run",
                    "-d",
                    "--name",
                    names[-1],
                    "--network",
                    network,
                    "--cpus",
                    cpu,
                    "--memory",
                    memory,
                    "--memory-swap",
                    memory,
                    "--env-file",
                    str(envfile),
                    "-v",
                    volume + ":/data",
                    *publish,
                    "ailab-final-backend:verify",
                    "uvicorn",
                    f"ailab_platform.{role}.main:app",
                    "--host",
                    "0.0.0.0",
                    "--port",
                    "8080",
                    "--no-proxy-headers",
                )
                if role == "api":
                    wait(f"http://127.0.0.1:{api_port}/readyz")
            names.append(prefix + "-portal")
            docker(
                "run",
                "-d",
                "--name",
                names[-1],
                "--network",
                network,
                "--cpus",
                "0.5",
                "--memory",
                "512m",
                "--memory-swap",
                "512m",
                "-e",
                f"AILAB_API_BASE=http://{prefix}-api:8080",
                "-p",
                f"127.0.0.1:{portal_port}:3000",
                "ailab-final-portal:verify",
                "node",
                "node_modules/next/dist/bin/next",
                "start",
                "--hostname",
                "0.0.0.0",
                "--port",
                "3000",
            )
            wait(endpoint + "/login")
            with httpx.Client(base_url=endpoint, timeout=90) as client:
                client.post("/api/auth/login", json={"username": "load-probe", "password": password}).raise_for_status()
                cookies = dict(client.cookies)
            source = root / "office.pptx"
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(source)
            with zipfile.ZipFile(source, "a") as archive:
                archive.writestr("padding.bin", os.urandom(20 * 1024**2), compress_type=zipfile.ZIP_STORED)

            def upload():
                started = time.monotonic()
                with httpx.Client(base_url=endpoint, cookies=cookies, timeout=90) as client, source.open("rb") as body:
                    response = client.post(
                        "/api/v1/files",
                        files={
                            "upload": ("office.pptx", body, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
                        },
                    )
                    response.raise_for_status()
                    fid = response.json()["id"]
                    deadline = time.monotonic() + 90
                    while time.monotonic() < deadline:
                        status = client.get(f"/api/v1/files/{fid}").json()["status"]
                        if status == "ready":
                            return {"file_id": fid, "seconds": time.monotonic() - started}
                        assert status != "rejected"
                        time.sleep(0.2)
                raise RuntimeError("Office parsing deadline")

            baseline_file = upload()["file_id"]

            def run(mode):
                started, first = time.monotonic(), None
                with httpx.Client(base_url=endpoint, cookies=cookies, timeout=90) as client:
                    response = client.post(
                        "/api/v1/runs", json={"prompt": "生成一个文件" if mode == "work" else "hello", "mode": mode, "model": "test-fast"}
                    )
                    response.raise_for_status()
                    rid = response.json()["id"]
                    with client.stream("GET", f"/api/v1/runs/{rid}/events") as stream:
                        stream.raise_for_status()
                        for line in stream.iter_lines():
                            if time.monotonic() - started > 90:
                                raise RuntimeError("SSE deadline")
                            if first is None and line == "event: agent.delta":
                                first = time.monotonic() - started
                    result = client.get(f"/api/v1/runs/{rid}").json()
                    assert result["status"] == "succeeded", result.get("error")
                    if mode == "work":
                        assert result["artifacts"]
                    return {"mode": mode, "seconds": time.monotonic() - started, "ttft_seconds": first}

            def download():
                size = 0
                with httpx.Client(base_url=endpoint, cookies=cookies, timeout=90) as client:
                    with client.stream("GET", f"/api/v1/files/{baseline_file}/content") as response:
                        response.raise_for_status()
                        for chunk in response.iter_bytes():
                            size += len(chunk)
                assert size == source.stat().st_size
                return size

            before = {role: counters(role) for role in (*limits, "portal")}
            started = time.monotonic()
            with concurrent.futures.ThreadPoolExecutor(max_workers=12) as pool:
                runs = [pool.submit(run, "work") for _ in range(2)] + [pool.submit(run, "chat") for _ in range(chats)]
                uploads = pool.submit(upload)
                downloads = [pool.submit(download) for _ in range(2)]
                results = [future.result(timeout=100) for future in runs]
                parsed = uploads.result(timeout=100)
                assert all(future.result(timeout=100) == source.stat().st_size for future in downloads)
            after = {role: counters(role) for role in (*limits, "portal")}
            for role, values in after.items():
                values["cpu_seconds"] -= before[role]["cpu_seconds"]
                assert values["oom_kill"] == 0
            latency = [result["seconds"] for result in results if result["mode"] == "chat"]
            ttft = [result["ttft_seconds"] for result in results if result["mode"] == "chat" and result["ttft_seconds"] is not None]
            assert len(ttft) == chats
            return {
                "chat_count": chats,
                "work_count": 2,
                "office_bytes": source.stat().st_size,
                "office_ready_seconds": parsed["seconds"],
                "chat_p95_seconds": statistics.quantiles(latency, n=100)[94],
                "ttft_p95_seconds": statistics.quantiles(ttft, n=100)[94],
                "wall_seconds": time.monotonic() - started,
                "services": after,
                "data_bytes": int(docker("exec", prefix + "-api", "du", "-sb", "/data").split()[0]),
                "scope": "five services and Redis; mock models/sandbox, stub scanner, local SQLite/storage; host load generator",
            }
    except BaseException:
        evidence = Path(tempfile.gettempdir()) / (prefix + "-failure-logs")
        evidence.mkdir(exist_ok=True)
        for name in names:
            result = subprocess.run(["docker", "--context", context, "logs", "--tail", "200", name], capture_output=True, timeout=20)
            (evidence / (name + ".log")).write_bytes(result.stdout + result.stderr)
        print("failure logs:", evidence, file=sys.stderr)
        raise
    finally:
        for name in reversed(names):
            subprocess.run(["docker", "--context", context, "rm", "-f", name], capture_output=True, timeout=30)
        for kind, name in (("volume", volume), ("network", network)):
            subprocess.run(["docker", "--context", context, kind, "rm", name], capture_output=True, timeout=30)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--context", default="colima-ailab-verify")
    parser.add_argument("--chats", type=int, default=20)
    args = parser.parse_args()
    if not 2 <= args.chats <= 100:
        parser.error("2 <= chats <= 100 required")
    print(json.dumps(measure(args.context, args.chats), indent=2))
